#!/usr/bin/env python3
"""
GitHub User Activity Report (GitHub App authentication, multi-org / multi-repo).

Requirements (Python 3.9+):
    pip install "requests>=2.31" "PyJWT[crypto]>=2.8" "openpyxl>=3.1"

GitHub App permissions (all read-only):
    Repository   : Metadata, Contents, Pull requests, Issues, Administration
    Organization : Members, Self-hosted runners

Usage:
    python gh_user_activity_report.py --config config.json
"""

import argparse
import csv
import ctypes
import hashlib
import json
import logging
import math
import os
import random
import re
import sys
import threading
import time
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from email.utils import parsedate_to_datetime
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional, Set, Tuple
from urllib.parse import urlparse

REQUIREMENTS = '"requests>=2.31" "PyJWT[crypto]>=2.8" "openpyxl>=3.1"'

try:
    import jwt
    import requests
    from openpyxl import Workbook
    from openpyxl.formatting.rule import DataBarRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from requests.adapters import HTTPAdapter
    from urllib3.util.retry import Retry
except ImportError as import_error:
    sys.exit(f"Missing dependency '{import_error.name}'. Install with: pip install {REQUIREMENTS}")

API_VERSION = "2022-11-28"
USER_AGENT = "gh-user-activity-report"
LOG = logging.getLogger("gh_user_activity")
IST = timezone(timedelta(hours=5, minutes=30), "IST")
DATE_FMT = "%d-%m-%Y"

ORG_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9-]{0,38}$")
REPO_RE = re.compile(r"^[A-Za-z0-9._-]{1,100}$")
PREFIX_RE = re.compile(r"^[A-Za-z0-9_.-]{1,64}$")
TRUE_VALUES = frozenset({"true", "yes", "y", "1"})
FALSE_VALUES = frozenset({"false", "no", "n", "0"})
REQUIRED_CSV_COLUMNS = frozenset({"org", "repo", "enabled"})
EMPTY_RESOURCE_STATUSES = frozenset({409, 410})  # empty repository / feature disabled
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

METRICS = (
    ("commits", "Commits"),
    ("prs_opened", "PRs Opened"),
    ("prs_merged", "PRs Merged"),
    ("reviews", "PR Reviews"),
    ("review_comments", "Review Comments"),
    ("issues_opened", "Issues Opened"),
    ("issue_comments", "Issue/PR Comments"),
)
METRIC_KEYS = tuple(key for key, _ in METRICS)
METRIC_LABELS = tuple(label for _, label in METRICS)
STATUSES = ("Active", "Moderate", "Low", "Inactive")
STATUS_ORDER = {status: rank for rank, status in enumerate(STATUSES)}

USER_HEADERS = (
    "Organization", "User", "Activity Status", "Active Days", "Total Activities",
    *METRIC_LABELS, "Average Commits per Active Month",
    "Repositories Contributed", "First Activity (IST)", "Last Activity (IST)",
    "Days Since Last Activity", "Org Member", "Teams", "Value Streams",
)
STATUS_IDX = USER_HEADERS.index("Activity Status")
TOTAL_IDX = USER_HEADERS.index("Total Activities")
METRIC_START_IDX = USER_HEADERS.index(METRIC_LABELS[0])

REPO_HEADERS = (
    "Organization", "Repository", "User", "Total Activities", *METRIC_LABELS,
    "Active Days", "Last Activity (IST)",
)
ORG_HEADERS = ("Organization", "Repositories Scanned", "Contributors", *STATUSES, "Total Activities")

RUNNER_HEADERS = (
    "Organization", "Runner ID", "Runner Name", "Scope", "Runner Group", "OS", "Arch", "Agent Version",
    "Provisioning Type", "Status", "Current State", "Labels", "Repos Served",
    "Total Jobs", "Successful", "Failed", "Cancelled", "Timed Out", "Success Rate (%)",
    "Total Compute (Hours)", "Avg Duration (Mins)", "Median Duration (Mins)",
    "Min Duration (Mins)", "Max Duration (Mins)",
    "Avg Queue Wait (s)", "P95 Queue Wait (s)", "SLA Breaches (>60s)",
    "First Active (IST)", "Last Active (IST)",
)
ORG_RUNNER_HEADERS = (
    "Organization", "Total Runners", "Org-Level", "Repo-Level",
    "Online", "Offline", "Availability (%)", "Currently Busy", "Ephemeral Count",
    "Active Repos", "Total Jobs Run", "Successful", "Failed", "Cancelled", "Timed Out",
    "Success Rate (%)", "Total Compute (Hours)",
    "Avg Queue Wait (s)", "P95 Queue Wait (s)", "SLA Breaches (>60s)", "SLA Breach Rate (%)",
)

NAVY = "1F3864"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
BAND_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
STATUS_STYLES = {
    "Active": ("C6EFCE", "006100"),
    "Moderate": ("FFEB9C", "9C5700"),
    "Low": ("FCE4D6", "C65911"),
    "Inactive": ("FFC7CE", "9C0006"),
}


class ConfigError(Exception):
    """Invalid configuration or input file."""


class GitHubAPIError(Exception):
    """Non-recoverable GitHub API response."""

    def __init__(self, status: int, message: str):
        super().__init__(message)
        self.status = status


# --------------------------------------------------------------------------- configuration

@dataclass(frozen=True)
class Config:
    api_url: str
    client_id: str
    private_key: str = field(repr=False)
    repos_csv: Path
    output_dir: Path
    file_prefix: str
    lookback_days: int
    active_min_days: int
    moderate_min_days: int
    include_archived: bool
    include_forks: bool
    exclude_bots: bool
    include_inactive_members: bool
    max_workers: int
    timeout: int
    max_retries: int
    log_level: str
    history_file: Optional[Path] = None
    value_stream_property: Tuple[str, ...] = ("VALUE_STREAM", "value_stream")
    audit_log_csv: Optional[Path] = None
    include_runner_metrics: bool = True
    max_runner_runs_per_repo: int = 50
    rate_limit_buffer: int = 25
    inter_request_delay_seconds: float = 0.05
    max_rate_limit_wait_seconds: int = 3600


@dataclass(frozen=True)
class RepoEntry:
    org: str
    repo: str
    enabled: bool


def _section(raw: Dict[str, Any], name: str) -> Dict[str, Any]:
    value = raw.get(name, {})
    if not isinstance(value, dict):
        raise ConfigError(f"'{name}' must be a JSON object")
    return value


def _req_str(section: Dict[str, Any], key: str, path: str) -> str:
    value = section.get(key)
    if not isinstance(value, str) or not value.strip():
        raise ConfigError(f"'{path}.{key}' is required and must be a non-empty string")
    return value.strip()


def _opt_int(section: Dict[str, Any], key: str, path: str, default: int, low: int, high: int) -> int:
    value = section.get(key, default)
    if isinstance(value, bool) or not isinstance(value, int) or not low <= value <= high:
        raise ConfigError(f"'{path}.{key}' must be an integer between {low} and {high}")
    return value


def _opt_bool(section: Dict[str, Any], key: str, path: str, default: bool) -> bool:
    value = section.get(key, default)
    if not isinstance(value, bool):
        raise ConfigError(f"'{path}.{key}' must be true or false")
    return value


def load_config(path: Path) -> Config:
    if not path.is_file():
        raise ConfigError(f"Config file not found: {path}")
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ConfigError(f"Invalid JSON in {path.name}: {exc}") from exc
    if not isinstance(raw, dict):
        raise ConfigError("Config root must be a JSON object")

    base_dir = path.resolve().parent
    github = _section(raw, "github")
    inputs = _section(raw, "input")
    output = _section(raw, "output")
    report = _section(raw, "report")
    runtime = _section(raw, "runtime")

    api_url = _req_str(github, "api_url", "github").rstrip("/")
    parsed = urlparse(api_url)
    if (parsed.scheme != "https" or not parsed.hostname or parsed.hostname == "github.com"
            or parsed.username or parsed.path not in ("", "/api/v3") or parsed.query or parsed.fragment):
        # Userinfo is dropped so embedded credentials are never echoed.
        received = f"{parsed.scheme}://{parsed.hostname or ''}{parsed.path}"
        raise ConfigError(f"'github.api_url' is '{received}' but must be the API root only: https://api.github.com, "
                          "https://api.<subdomain>.ghe.com or https://<ghes-host>/api/v3")

    key_path = base_dir / _req_str(github, "private_key_path", "github")
    if not key_path.is_file():
        raise ConfigError(f"GitHub App private key not found: {key_path}")
    private_key = key_path.read_text(encoding="ascii")
    if "PRIVATE KEY" not in private_key:
        raise ConfigError(f"File is not a PEM private key: {key_path}")

    file_prefix = output.get("file_prefix", "github_user_activity")
    if not isinstance(file_prefix, str) or not PREFIX_RE.fullmatch(file_prefix):
        raise ConfigError("'output.file_prefix' may only contain letters, digits, '.', '_' and '-'")

    active_min = _opt_int(report, "active_min_days", "report", 8, 3, 3650)
    moderate_min = _opt_int(report, "moderate_min_days", "report", 3, 2, 3650)
    if moderate_min >= active_min:
        raise ConfigError("'report.moderate_min_days' must be lower than 'report.active_min_days'")

    log_level = str(runtime.get("log_level", "INFO")).upper()
    if log_level not in {"DEBUG", "INFO", "WARNING", "ERROR"}:
        raise ConfigError("'runtime.log_level' must be DEBUG, INFO, WARNING or ERROR")
    history_file = report.get("history_file")
    if history_file is not None and (not isinstance(history_file, str) or not history_file.strip()):
        raise ConfigError("'report.history_file' must be a non-empty path string")
    audit_log_csv = inputs.get("audit_log_csv")
    if audit_log_csv is not None and (not isinstance(audit_log_csv, str) or not audit_log_csv.strip()):
        raise ConfigError("'input.audit_log_csv' must be a non-empty path string or null")
    audit_log_path = base_dir / audit_log_csv if audit_log_csv is not None else None
    if audit_log_path is not None and not audit_log_path.is_file():
        raise ConfigError(f"Audit-log CSV not found: {audit_log_path}")
    property_names = report.get("value_stream_property", ["VALUE_STREAM", "value_stream"])
    if isinstance(property_names, str):
        property_names = [property_names]
    if (not isinstance(property_names, list) or not property_names
            or any(not isinstance(name, str) or not name.strip() or len(name) > 120
                   or any(ord(char) < 32 for char in name) for name in property_names)):
        raise ConfigError("'report.value_stream_property' must be a property name or a non-empty list of names")

    inter_request_delay = float(runtime.get("inter_request_delay_seconds", 0.05))
    if not (0.0 <= inter_request_delay <= 5.0):
        raise ConfigError("'runtime.inter_request_delay_seconds' must be between 0.0 and 5.0")

    return Config(
        api_url=api_url,
        client_id=_req_str(github, "app_client_id", "github"),
        private_key=private_key,
        repos_csv=base_dir / _req_str(inputs, "repos_csv", "input"),
        output_dir=base_dir / str(output.get("directory", "reports")),
        file_prefix=file_prefix,
        lookback_days=_opt_int(report, "lookback_days", "report", 90, 1, 3650),
        active_min_days=active_min,
        moderate_min_days=moderate_min,
        include_archived=_opt_bool(report, "include_archived_repos", "report", False),
        include_forks=_opt_bool(report, "include_forked_repos", "report", True),
        exclude_bots=_opt_bool(report, "exclude_bots", "report", True),
        include_inactive_members=_opt_bool(report, "include_inactive_org_members", "report", True),
        max_workers=_opt_int(runtime, "max_workers", "runtime", 4, 1, 16),
        timeout=_opt_int(runtime, "request_timeout_seconds", "runtime", 30, 5, 300),
        max_retries=_opt_int(runtime, "max_retries", "runtime", 5, 0, 10),
        log_level=log_level,
        history_file=base_dir / history_file if history_file is not None else None,
        value_stream_property=tuple(dict.fromkeys(name.strip() for name in property_names)),
        audit_log_csv=audit_log_path,
        include_runner_metrics=_opt_bool(report, "include_runner_metrics", "report", True),
        max_runner_runs_per_repo=_opt_int(report, "max_runner_runs_per_repo", "report", 50, 1, 1000),
        rate_limit_buffer=_opt_int(runtime, "rate_limit_buffer", "runtime", 25, 0, 500),
        inter_request_delay_seconds=inter_request_delay,
        max_rate_limit_wait_seconds=_opt_int(runtime, "max_rate_limit_wait_seconds", "runtime", 3600, 60, 86400),
    )


def load_repo_entries(path: Path) -> List[RepoEntry]:
    if not path.is_file():
        raise ConfigError(f"Repository CSV not found: {path}")

    entries: List[RepoEntry] = []
    seen: Dict[Tuple[str, str], bool] = {}
    errors: List[str] = []
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        columns = {(name or "").strip().lower() for name in reader.fieldnames or []}
        missing = REQUIRED_CSV_COLUMNS - columns
        if missing:
            raise ConfigError(f"{path.name} is missing required column(s): {', '.join(sorted(missing))}")

        for line_no, raw in enumerate(reader, start=2):
            row = {k.strip().lower(): v.strip() for k, v in raw.items() if isinstance(k, str) and isinstance(v, str)}
            if not any(row.values()):
                continue
            org, repo, flag = row.get("org", ""), row.get("repo", ""), row.get("enabled", "").lower()

            problems = []
            if not ORG_RE.fullmatch(org):
                problems.append(f"invalid org '{org}' (use the login only, e.g. 'my-org', not a URL)")
            if repo != "*" and (not REPO_RE.fullmatch(repo) or repo in {".", ".."}):
                problems.append(f"invalid repo '{repo}' (use * for all repositories)")
            if flag not in TRUE_VALUES | FALSE_VALUES:
                problems.append("enabled must be true/false/yes/no/1/0")
            if problems:
                errors.append(f"line {line_no}: {'; '.join(problems)}")
                continue

            enabled = flag in TRUE_VALUES
            key = (org.lower(), repo.lower())
            if key in seen:
                if seen[key] != enabled:
                    errors.append(f"line {line_no}: conflicting enabled values for {org}/{repo}")
                continue
            seen[key] = enabled
            entries.append(RepoEntry(org, repo, enabled))

    if errors:
        raise ConfigError(f"{path.name} validation failed:\n  " + "\n  ".join(errors))
    if not any(entry.enabled for entry in entries):
        raise ConfigError(f"{path.name} contains no enabled rows")
    return entries


# --------------------------------------------------------------------------- GitHub client

class GitHubClient:
    """Thread-safe REST client authenticating as a GitHub App installation per organization."""

    def __init__(self, cfg: Config):
        self._cfg = cfg
        self._host = urlparse(cfg.api_url).netloc
        self._local = threading.local()
        self._lock = threading.Lock()
        self._installations: Dict[str, int] = {}
        self._tokens: Dict[str, Tuple[str, float]] = {}
        self._clock_offset: Optional[int] = None

        # Cross-thread rate-limiting & pacing coordination
        self._rate_limit_lock = threading.Lock()
        self._global_pause_until: float = 0.0
        self._app_pause_until: float = 0.0
        self._app_remaining: Optional[int] = None
        self._app_reset: Optional[float] = None
        self._org_pause_until: Dict[str, float] = {}
        self._org_remaining: Dict[str, int] = {}
        self._org_reset: Dict[str, float] = {}
        self._last_request_time: float = 0.0
        self._min_delay: float = cfg.inter_request_delay_seconds
        self._rate_limit_buffer: int = cfg.rate_limit_buffer
        self._max_rate_limit_wait: int = cfg.max_rate_limit_wait_seconds

    def _http(self) -> requests.Session:
        session = getattr(self._local, "session", None)
        if session is None:
            retry = Retry(
                total=self._cfg.max_retries,
                backoff_factor=1.5,
                status_forcelist=(500, 502, 503, 504),
                allowed_methods=frozenset({"GET", "POST"}),
                raise_on_status=False,
            )
            session = requests.Session()
            session.mount("https://", HTTPAdapter(max_retries=retry))
            session.headers.update({
                "Accept": "application/vnd.github+json",
                "X-GitHub-Api-Version": API_VERSION,
                "User-Agent": USER_AGENT,
            })
            self._local.session = session
        return session

    def _server_now(self) -> int:
        """Current time per GitHub's clock; JWTs are rejected if the local clock drifts."""
        if self._clock_offset is None:
            server_date = self._http().head(self._cfg.api_url, timeout=self._cfg.timeout).headers.get("Date")
            self._clock_offset = (int(parsedate_to_datetime(server_date).timestamp() - time.time())
                                  if server_date else 0)
        return int(time.time()) + self._clock_offset

    def _app_jwt(self) -> str:
        now = self._server_now()
        payload = {"iat": now - 60, "exp": now + 540, "iss": self._cfg.client_id}
        return jwt.encode(payload, self._cfg.private_key, algorithm="RS256")

    def authenticate(self, org: str) -> str:
        """Return a cached installation token for the org, refreshing 5 minutes before expiry."""
        key = org.lower()
        with self._lock:
            cached = self._tokens.get(key)
            if cached and cached[1] - self._server_now() > 300:
                return cached[0]

            headers = {"Authorization": f"Bearer {self._app_jwt()}"}
            if key not in self._installations:
                self._installations = self._list_installations(headers)
            if key not in self._installations:
                installed_on = ", ".join(sorted(self._installations)) or "none"
                raise GitHubAPIError(404, f"GitHub App is not installed on '{org}' at {self._host} "
                                          f"(App installations found: {installed_on})")

            resp = self._send("POST", f"/app/installations/{self._installations[key]}/access_tokens",
                              headers, is_app_auth=True)
            self._raise_for_status(resp)
            data = resp.json()
            expires_at = parse_ts(data["expires_at"]).timestamp()
            self._tokens[key] = (data["token"], expires_at)
            LOG.debug("Issued installation token for %s", org)
            return data["token"]

    def _list_installations(self, headers: Dict[str, str]) -> Dict[str, int]:
        """Map installed account logins (lower-case) to installation IDs."""
        installations: Dict[str, int] = {}
        url: Optional[str] = "/app/installations"
        query: Optional[Dict[str, Any]] = {"per_page": 100}
        while url:
            resp = self._send("GET", url, headers, query, is_app_auth=True)
            if resp.status_code == 401:
                raise GitHubAPIError(401, "GitHub App authentication failed - verify client ID, "
                                          "private key, api_url and system clock")
            self._raise_for_status(resp)
            for item in resp.json():
                installations[item["account"]["login"].lower()] = item["id"]
            url, query = self._next_url(resp), None
        return installations

    def _next_url(self, resp: requests.Response) -> Optional[str]:
        url = resp.links.get("next", {}).get("url")
        if url and urlparse(url).netloc != self._host:
            raise GitHubAPIError(0, f"Refusing to follow pagination link to foreign host: {urlparse(url).netloc}")
        return url

    def _coordinate_rate_limit(self, org: Optional[str], is_app_auth: bool = False) -> None:
        """Pacing and proactive/reactive rate limit barrier coordinated across all worker threads."""
        org_key = org.lower() if org else None
        pacing_wait = 0.0
        while True:
            with self._rate_limit_lock:
                now = time.time()
                pause_wait = 0.0

                # 1. Global pause (secondary rate limit, 429, Retry-After)
                if self._global_pause_until > now:
                    pause_wait = max(pause_wait, self._global_pause_until - now)

                # 2. App-level JWT rate limit (applies to /app/* endpoints using JWT bearer auth: 60 req/hr)
                if is_app_auth:
                    if self._app_pause_until > now:
                        pause_wait = max(pause_wait, self._app_pause_until - now)
                    elif pause_wait == 0.0 and self._app_remaining is not None and self._app_remaining <= 5:
                        reset_at = self._app_reset or (now + 60.0)
                        needed = max(0.0, reset_at - now) + 2.0
                        self._app_pause_until = now + needed
                        pause_wait = needed
                        LOG.warning("Proactive App-level rate limit pause: %d JWT requests remaining; waiting %.1fs until reset",
                                    self._app_remaining, needed)

                # 3. Org-specific pause (primary installation rate limit reset for this installation: 5,000-12,500 req/hr)
                if not is_app_auth and org_key:
                    if self._org_pause_until.get(org_key, 0.0) > now:
                        pause_wait = max(pause_wait, self._org_pause_until[org_key] - now)
                    elif pause_wait == 0.0:
                        remaining = self._org_remaining.get(org_key)
                        if remaining is not None and remaining <= self._rate_limit_buffer:
                            reset_at = self._org_reset.get(org_key, now + 60.0)
                            needed = max(0.0, reset_at - now) + 2.0
                            if needed > self._max_rate_limit_wait:
                                LOG.error("Rate limit reset for %s is %.0fs away, exceeding max wait %ds",
                                          org, needed, self._max_rate_limit_wait)
                            else:
                                self._org_pause_until[org_key] = now + needed
                                pause_wait = needed
                                LOG.warning("Proactive installation rate limit pause for %s: %d requests remaining; waiting %.1fs until reset",
                                            org, remaining, needed)

                if pause_wait > 0.0:
                    pass
                else:
                    # 4. Atomic slot reservation for inter-request pacing
                    if self._min_delay > 0.0:
                        target = max(now, self._last_request_time + self._min_delay)
                        self._last_request_time = target
                        pacing_wait = target - now
                    else:
                        pacing_wait = 0.0
                    break

            time.sleep(min(pause_wait, 5.0))

        if pacing_wait > 0.0:
            time.sleep(pacing_wait)

    def get(self, org: str, path: str, params: Optional[Dict[str, Any]] = None) -> Any:
        resp = self._send("GET", path, {"Authorization": f"Bearer {self.authenticate(org)}"}, params, org=org)
        self._raise_for_status(resp)
        return resp.json()

    def get_response(self, org: str, path: str, params: Optional[Dict[str, Any]] = None) -> requests.Response:
        return self._send("GET", path, {"Authorization": f"Bearer {self.authenticate(org)}"}, params, org=org)

    def paginate(self, org: str, path: str, params: Optional[Dict[str, Any]] = None) -> Iterator[Dict[str, Any]]:
        url: Optional[str] = path
        query: Optional[Dict[str, Any]] = {**(params or {}), "per_page": 100}
        while url:
            resp = self._send("GET", url, {"Authorization": f"Bearer {self.authenticate(org)}"}, query, org=org)
            if resp.status_code in EMPTY_RESOURCE_STATUSES:
                return
            self._raise_for_status(resp)
            yield from resp.json()
            url, query = self._next_url(resp), None

    def paginate_enveloped(self, org: str, path: str, envelope_key: str,
                           params: Optional[Dict[str, Any]] = None) -> Iterator[Dict[str, Any]]:
        url: Optional[str] = path
        query: Optional[Dict[str, Any]] = {**(params or {}), "per_page": 100}
        while url:
            resp = self._send("GET", url, {"Authorization": f"Bearer {self.authenticate(org)}"}, query, org=org)
            if resp.status_code in EMPTY_RESOURCE_STATUSES:
                return
            self._raise_for_status(resp)
            data = resp.json()
            if isinstance(data, dict):
                items = data.get(envelope_key, [])
                if isinstance(items, list):
                    yield from items
            elif isinstance(data, list):
                yield from data
            url, query = self._next_url(resp), None

    def _send(self, method: str, url: str, headers: Dict[str, str],
              params: Optional[Dict[str, Any]] = None,
              org: Optional[str] = None,
              is_app_auth: bool = False) -> requests.Response:
        if not url.startswith("https://"):
            url = f"{self._cfg.api_url}{url}"
        org_key = org.lower() if org else None

        for attempt in range(self._cfg.max_retries + 1):
            self._coordinate_rate_limit(org, is_app_auth=is_app_auth)
            resp = self._http().request(method, url, headers=headers, params=params, timeout=self._cfg.timeout)

            # Record rate limit headers for proactive tracking
            rem = resp.headers.get("X-RateLimit-Remaining")
            rst = resp.headers.get("X-RateLimit-Reset")
            with self._rate_limit_lock:
                if is_app_auth:
                    if rem is not None and rem.isdigit():
                        self._app_remaining = int(rem)
                    if rst is not None and rst.isdigit():
                        self._app_reset = float(rst)
                elif org_key:
                    if rem is not None and rem.isdigit():
                        self._org_remaining[org_key] = int(rem)
                    if rst is not None and rst.isdigit():
                        self._org_reset[org_key] = float(rst)

            wait = self._rate_limit_wait(resp)
            if wait is None:
                return resp
            if attempt == self._cfg.max_retries:
                break

            is_secondary = (resp.status_code == 429
                            or "secondary rate limit" in resp.text.lower()
                            or "abuse detection" in resp.text.lower()
                            or "please wait" in resp.text.lower())
            jitter = random.uniform(0.5, 2.0)
            actual_wait = wait + jitter

            with self._rate_limit_lock:
                now = time.time()
                if is_secondary or (not org_key and not is_app_auth):
                    backoff_wait = max(actual_wait, (2 ** attempt) * 15.0 + jitter)
                    self._global_pause_until = max(self._global_pause_until, now + backoff_wait)
                    LOG.warning("Secondary / concurrency rate limit hit (HTTP %s) on %s; pausing all requests for %.1fs (attempt %d/%d)",
                                resp.status_code, urlparse(url).path, backoff_wait, attempt + 1, self._cfg.max_retries)
                    actual_wait = backoff_wait
                elif is_app_auth:
                    self._app_pause_until = max(self._app_pause_until, now + actual_wait)
                    LOG.warning("App-level JWT rate limit hit (HTTP %s) on %s; pausing App token requests for %.1fs (attempt %d/%d)",
                                resp.status_code, urlparse(url).path, actual_wait, attempt + 1, self._cfg.max_retries)
                else:
                    self._org_pause_until[org_key] = max(self._org_pause_until.get(org_key, 0.0), now + actual_wait)
                    LOG.warning("Installation rate limit hit for %s (HTTP %s) on %s; pausing org for %.1fs (attempt %d/%d)",
                                org, resp.status_code, urlparse(url).path, actual_wait, attempt + 1, self._cfg.max_retries)

            time.sleep(actual_wait)

        raise GitHubAPIError(resp.status_code, f"Rate limit retries exhausted for {urlparse(url).path}")

    @staticmethod
    def _rate_limit_wait(resp: requests.Response) -> Optional[int]:
        if resp.status_code not in (403, 429):
            return None
        retry_after = resp.headers.get("Retry-After")
        if retry_after and retry_after.isdigit():
            return int(retry_after) + 1
        text = resp.text.lower()
        if "secondary rate limit" in text or "abuse detection" in text or "please wait" in text:
            return 60
        if resp.headers.get("X-RateLimit-Remaining") == "0":
            reset = int(resp.headers.get("X-RateLimit-Reset", "0"))
            try:
                server_date = parsedate_to_datetime(resp.headers["Date"])
                now = server_date.timestamp()
            except (KeyError, TypeError, ValueError, OverflowError):
                now = time.time()
            wait = reset - int(now)
            return wait + 2 if wait > 0 else 60
        if "rate limit" in text:
            reset = int(resp.headers.get("X-RateLimit-Reset", "0"))
            if reset > 0:
                try:
                    server_date = parsedate_to_datetime(resp.headers["Date"])
                    now = server_date.timestamp()
                except (KeyError, TypeError, ValueError, OverflowError):
                    now = time.time()
                wait = reset - int(now)
                if wait > 0:
                    return wait + 2
            return 60
        if resp.status_code == 429:
            return 60
        return None
    @staticmethod
    def _raise_for_status(resp: requests.Response) -> None:
        if resp.status_code < 400:
            return
        try:
            body = resp.json()
            message = body.get("message", "") if isinstance(body, dict) else ""
        except ValueError:
            message = resp.text[:200]
        raise GitHubAPIError(resp.status_code, f"HTTP {resp.status_code} on {urlparse(resp.url).path}: {message}")


# --------------------------------------------------------------------------- data collection

@dataclass
class Activity:
    counts: Dict[str, int] = field(default_factory=lambda: dict.fromkeys(METRIC_KEYS, 0))
    days: Set[date] = field(default_factory=set)
    first: Optional[datetime] = None
    last: Optional[datetime] = None

    @property
    def total(self) -> int:
        return sum(self.counts.values())

    def add(self, metric: str, when: datetime) -> None:
        self.counts[metric] += 1
        self.days.add(when.astimezone(IST).date())
        self._extend(when, when)

    def merge(self, other: "Activity") -> None:
        for metric, count in other.counts.items():
            self.counts[metric] += count
        self.days |= other.days
        if other.first and other.last:
            self._extend(other.first, other.last)

    def _extend(self, first: datetime, last: datetime) -> None:
        self.first = first if self.first is None else min(self.first, first)
        self.last = last if self.last is None else max(self.last, last)


def parse_ts(value: str) -> datetime:
    return datetime.fromisoformat(value.replace("Z", "+00:00")).astimezone(timezone.utc)


def parse_audit_timestamp(value: str) -> datetime:
    try:
        return parse_ts(value)
    except ValueError:
        try:
            epoch_value = int(value)
            divisor = 1000 if abs(epoch_value) >= 100_000_000_000 else 1
            return datetime.fromtimestamp(epoch_value / divisor, timezone.utc)
        except (ValueError, OverflowError, OSError) as exc:
            raise ValueError(f"invalid audit-log timestamp '{value}'") from exc


def load_onboarded_users(path: Path, since: datetime, end: datetime) -> Dict[str, Set[str]]:
    users_by_org: Dict[str, Set[str]] = defaultdict(set)
    try:
        with path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            columns = {str(name).strip().casefold() for name in reader.fieldnames or []}
            missing = {"org", "action", "user"} - columns
            timestamp_column = "created_at" if "created_at" in columns else "@timestamp"
            if missing or timestamp_column not in columns:
                required = sorted(missing | {"created_at or @timestamp"} if timestamp_column not in columns else missing)
                raise ConfigError(f"{path.name} is missing audit-log column(s): {', '.join(required)}")

            for line_no, raw in enumerate(reader, start=2):
                row = {str(key).strip().casefold(): (value or "").strip()
                       for key, value in raw.items() if key is not None}
                if row.get("action", "").casefold() != "org.add_member":
                    continue
                org, user = row.get("org", ""), row.get("user", "")
                if not org or not user:
                    raise ConfigError(f"{path.name} line {line_no}: org.add_member event is missing org or user")
                try:
                    occurred_at = parse_audit_timestamp(row.get(timestamp_column, ""))
                except ValueError as exc:
                    raise ConfigError(f"{path.name} line {line_no}: {exc}") from exc
                if since <= occurred_at <= end:
                    users_by_org[org.casefold()].add(user.casefold())
    except OSError as exc:
        raise ConfigError(f"Could not read audit-log CSV {path.name}: {exc}") from exc
    return users_by_org


def fmt_timestamp(value: Optional[datetime]) -> Optional[str]:
    return value.astimezone(IST).strftime(f"{DATE_FMT} %H:%M") if value else None


def is_bot(user: Dict[str, Any]) -> bool:
    return user.get("type") == "Bot" or str(user.get("login", "")).endswith("[bot]")


@dataclass
class RepositoryReport:
    metadata: Dict[str, Any] = field(default_factory=dict)
    value_streams: Tuple[str, ...] = ("Unknown",)
    users: Dict[str, Activity] = field(default_factory=dict)
    events: List[Tuple[Optional[str], str, datetime]] = field(default_factory=list)
    protected: str = "Unknown"
    protected_branches: Tuple[str, ...] = ()
    codeowners: str = "Unknown"
    codeowners_path: str = ""
    review_hours: List[float] = field(default_factory=list)
    cohort_opened: int = 0
    cohort_merged: int = 0
    activity_complete: bool = False
    errors: List[str] = field(default_factory=list)


def collect_repo(client: GitHubClient, org: str, repo: str, since: datetime,
                 exclude_bots: bool, end: Optional[datetime] = None,
                 report: Optional[RepositoryReport] = None) -> Dict[str, Activity]:
    activity: Dict[str, Activity] = defaultdict(Activity)
    base = f"/repos/{org}/{repo}"
    end = end or datetime.now(timezone.utc)
    month_start = end.astimezone(IST).replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    collected_since = min(since, end - timedelta(days=30), month_start) if report else since
    since_param = {"since": collected_since.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")}

    def record(user: Optional[Dict[str, Any]], metric: str, timestamp: Optional[str]) -> None:
        if not timestamp or (user and exclude_bots and is_bot(user)):
            return
        when = parse_ts(timestamp)
        login = user.get("login") if user else None
        if collected_since <= when <= end and report is not None:
            report.events.append((login.lower() if login else None, metric, when))
        if since <= when <= end and login:
            activity[login].add(metric, when)

    for commit in client.paginate(org, f"{base}/commits", {
            **since_param, "until": end.strftime("%Y-%m-%dT%H:%M:%SZ")}):
        record(commit.get("author"), "commits", commit["commit"]["author"]["date"])

    pull_requests_enabled = report is None or report.metadata.get("has_pull_requests", True)
    if pull_requests_enabled:
        pulls = client.paginate(org, f"{base}/pulls", {"state": "all", "sort": "updated", "direction": "desc"})
        for pr in pulls:
            if parse_ts(pr["updated_at"]) < collected_since:
                break
            record(pr.get("user"), "prs_opened", pr["created_at"])
            record(pr.get("user"), "prs_merged", pr.get("merged_at"))
            created = parse_ts(pr["created_at"])
            eligible = not (exclude_bots and is_bot(pr.get("user") or {}))
            if report is not None and eligible and since <= created <= end:
                report.cohort_opened += 1
                if pr.get("merged_at") and parse_ts(pr["merged_at"]) <= end:
                    report.cohort_merged += 1
            first_review: Optional[datetime] = None
            for review in client.paginate(org, f"{base}/pulls/{pr['number']}/reviews"):
                reviewer = review.get("user") or {}
                if review.get("state") == "PENDING" or not review.get("submitted_at"):
                    continue
                record(reviewer, "reviews", review["submitted_at"])
                submitted = parse_ts(review["submitted_at"])
                if (reviewer.get("login") and reviewer["login"].lower() != (pr.get("user") or {}).get("login", "").lower()
                        and not (exclude_bots and is_bot(reviewer)) and created <= submitted <= end):
                    first_review = min(first_review, submitted) if first_review else submitted
            if report is not None and first_review and since <= first_review <= end:
                report.review_hours.append((first_review - created).total_seconds() / 3600)

    for issue in client.paginate(org, f"{base}/issues", {"state": "all", **since_param}):
        if "pull_request" not in issue:
            record(issue.get("user"), "issues_opened", issue["created_at"])

    for comment in client.paginate(org, f"{base}/issues/comments", since_param):
        record(comment.get("user"), "issue_comments", comment["created_at"])

    if pull_requests_enabled:
        for comment in client.paginate(org, f"{base}/pulls/comments", since_param):
            record(comment.get("user"), "review_comments", comment["created_at"])

    return dict(activity)


def property_value_streams(properties: Any, property_names: Any = ("VALUE_STREAM", "value_stream")) -> Tuple[str, ...]:
    if isinstance(properties, dict):
        properties = [{"property_name": name, "value": value} for name, value in properties.items()]
    if not isinstance(properties, list) or any(not isinstance(item, dict) for item in properties):
        raise ValueError("Invalid repository custom-property response")
    if isinstance(property_names, str):
        property_names = [property_names]
    if (not isinstance(property_names, (list, tuple)) or not property_names
            or any(not isinstance(name, str) or not name.strip() for name in property_names)):
        raise ValueError("Value-stream property names must be a non-empty list of strings")
    normalized_names = {re.sub(r"[\s_-]", "", name).casefold() for name in property_names}
    matching = [item for item in properties if isinstance(item.get("property_name"), str)
                and re.sub(r"[\s_-]", "", item["property_name"]).casefold() in normalized_names]
    if not matching:
        return ("Unmapped",)
    values: Set[str] = set()
    for prop in matching:
        value = prop.get("value")
        if value is None:
            continue
        property_values = [value] if isinstance(value, str) else value
        if not isinstance(property_values, list) or any(not isinstance(item, str) for item in property_values):
            raise ValueError("Value-stream property must be a string, list of strings, or null")
        values.update(part.strip() for item in property_values for part in item.split(";") if part.strip())
    mapped_values = {value for value in values if value.casefold() != "unmapped"}
    return tuple(sorted(mapped_values, key=str.casefold)) or ("Unmapped",)


def collect_repository(client: GitHubClient, org: str, repo: str, since: datetime,
                       end: datetime, exclude_bots: bool,
                       property_names: Tuple[str, ...] = ("VALUE_STREAM", "value_stream")) -> RepositoryReport:
    LOG.info("Scanning repository %s/%s...", org, repo)
    report = RepositoryReport()
    base = f"/repos/{org}/{repo}"
    report.metadata = client.get(org, base)
    try:
        report.value_streams = property_value_streams(client.get(org, f"{base}/properties/values"), property_names)
    except (GitHubAPIError, requests.RequestException, ValueError) as exc:
        report.errors.append(f"Value-stream custom property: {exc}")
    if report.value_streams == ("Unmapped",):
        metadata_properties = report.metadata.get("custom_properties")
        if isinstance(metadata_properties, dict):
            try:
                metadata_streams = property_value_streams(metadata_properties, property_names)
            except ValueError:
                metadata_streams = ("Unmapped",)
            if metadata_streams != ("Unmapped",):
                report.value_streams = metadata_streams
        if report.value_streams == ("Unmapped",):
            LOG.debug("No configured value-stream property found for %s/%s; available custom property names: %s",
                      org, repo, sorted(metadata_properties) if isinstance(metadata_properties, dict) else "not included")
    try:
        report.users = collect_repo(client, org, repo, since, exclude_bots, end, report)
        report.activity_complete = True
    except (GitHubAPIError, requests.RequestException, KeyError, ValueError) as exc:
        report.errors.append(f"Activity: {exc}")
        report.events.clear()
        report.review_hours.clear()
        report.cohort_opened = report.cohort_merged = 0
    try:
        report.protected_branches = tuple(sorted(
            (branch["name"] for branch in client.paginate(org, f"{base}/branches", {"protected": "true"})),
            key=str.casefold))
        report.protected = "Yes" if report.protected_branches else "No"
    except (GitHubAPIError, requests.RequestException) as exc:
        report.errors.append(f"Branch protection: {exc}")
    try:
        branch = report.metadata.get("default_branch")
        if not branch:
            report.codeowners = "No"
        else:
            for path in (".github/CODEOWNERS", "CODEOWNERS", "docs/CODEOWNERS"):
                try:
                    content = client.get(org, f"{base}/contents/{path}", {"ref": branch})
                except GitHubAPIError as exc:
                    if exc.status == 404:
                        continue
                    raise
                if isinstance(content, dict) and content.get("type") == "file":
                    report.codeowners, report.codeowners_path = "Yes", path
                    break
            else:
                report.codeowners = "No"
    except (GitHubAPIError, requests.RequestException) as exc:
        report.errors.append(f"CODEOWNERS: {exc}")
    return report


def resolve_targets(client: GitHubClient, entries: List[RepoEntry],
                    cfg: Config) -> Tuple[List[Tuple[str, str]], List[Tuple[str, str]]]:
    by_org: Dict[str, List[RepoEntry]] = defaultdict(list)
    for entry in entries:
        by_org[entry.org.lower()].append(entry)

    targets: List[Tuple[str, str]] = []
    failures: List[Tuple[str, str]] = []
    for rows in by_org.values():
        org = rows[0].org
        wildcard = [row for row in rows if row.repo == "*"]
        if any(not row.enabled for row in wildcard):
            LOG.info("Skipping organization %s (disabled in CSV)", org)
            continue
        include = {row.repo.lower(): row.repo for row in rows if row.repo != "*" and row.enabled}
        exclude = {row.repo.lower() for row in rows if row.repo != "*" and not row.enabled}
        if not wildcard and not include:
            continue
        try:
            client.authenticate(org)
            if wildcard:
                for repo in client.paginate(org, f"/orgs/{org}/repos", {"type": "all"}):
                    if (repo["archived"] and not cfg.include_archived) or (repo["fork"] and not cfg.include_forks):
                        continue
                    include.setdefault(repo["name"].lower(), repo["name"])
        except (GitHubAPIError, requests.RequestException) as exc:
            LOG.error("Organization %s skipped: %s", org, exc)
            failures.append((org, str(exc)))
            continue
        selected = [(org, name) for key, name in sorted(include.items()) if key not in exclude]
        LOG.info("Organization %s: %d repositories selected", org, len(selected))
        targets.extend(selected)
    return targets, failures


def fetch_members(client: GitHubClient, orgs: List[str], exclude_bots: bool,
                  failures: List[Tuple[str, str]]) -> Dict[str, Set[str]]:
    members: Dict[str, Set[str]] = {}
    for org in orgs:
        try:
            members[org] = {
                m["login"] for m in client.paginate(org, f"/orgs/{org}/members")
                if not (exclude_bots and is_bot(m))
            }
            LOG.info("Organization %s: %d members", org, len(members[org]))
        except (GitHubAPIError, requests.RequestException) as exc:
            LOG.warning("Could not list members of %s: %s", org, exc)
            failures.append((f"{org} (member list)", str(exc)))
    return members


def fetch_team_access(client: GitHubClient, orgs: List[str], reports: Dict[Tuple[str, str], RepositoryReport],
                      exclude_bots: bool, failures: List[Tuple[str, str]],
                      property_names: Tuple[str, ...] = ("VALUE_STREAM", "value_stream")) -> Tuple[
                          Dict[Tuple[str, str], Set[str]], Dict[Tuple[str, str], Set[str]], Dict[str, bool]]:
    teams_by_user: Dict[Tuple[str, str], Set[str]] = defaultdict(set)
    streams_by_user: Dict[Tuple[str, str], Set[str]] = defaultdict(set)
    org_team_complete: Dict[str, bool] = {}

    for org in orgs:
        org_team_complete[org] = True
        selected_repo_streams = {
            (owner.casefold(), repo.casefold()): result.value_streams
            for (owner, repo), result in reports.items() if owner.casefold() == org.casefold()
        }
        try:
            teams = list(client.paginate(org, f"/orgs/{org}/teams"))
        except (GitHubAPIError, requests.RequestException) as exc:
            LOG.warning("Could not list teams for %s: %s", org, exc)
            failures.append((f"{org} (team list)", str(exc)))
            org_team_complete[org] = False
            continue

        team_repository_refs: Set[Tuple[str, str]] = set()
        matched_repository_refs: Set[Tuple[str, str]] = set()
        mapped_repository_refs: Set[Tuple[str, str]] = set()
        team_members_by_slug: Dict[str, Tuple[str, Set[str]]] = {}
        for team in teams:
            slug = team.get("slug")
            name = team.get("name") or slug
            if not slug or not name:
                failures.append((f"{org} (team entry)", "GitHub returned a team without a name or slug"))
                org_team_complete[org] = False
                continue
            try:
                team_members = list(client.paginate(org, f"/orgs/{org}/teams/{slug}/members"))
            except (GitHubAPIError, requests.RequestException) as exc:
                LOG.warning("Could not list members for team %s/%s: %s", org, slug, exc)
                failures.append((f"{org} (team {slug} members)", str(exc)))
                org_team_complete[org] = False
                continue
            member_logins = {
                member["login"] for member in team_members
                if member.get("login") and not (exclude_bots and is_bot(member))
            }
            team_members_by_slug[slug.casefold()] = (name, member_logins)
            try:
                team_repositories = list(client.paginate(org, f"/orgs/{org}/teams/{slug}/repos"))
            except (GitHubAPIError, requests.RequestException) as exc:
                LOG.warning("Could not list repositories for team %s/%s: %s", org, slug, exc)
                failures.append((f"{org} (team {slug} repositories)", str(exc)))
                org_team_complete[org] = False
                team_repositories = []

            team_streams: Set[str] = set()
            for repository in team_repositories:
                repo_name = repository.get("name") or repository.get("full_name", "").rsplit("/", 1)[-1]
                full_name = repository.get("full_name", "")
                owner_name = repository.get("owner", {}).get("login") if isinstance(repository.get("owner"), dict) else None
                owner_name = owner_name or (full_name.split("/", 1)[0] if "/" in full_name else org)
                repo_key = (owner_name.casefold(), repo_name.casefold())
                team_repository_refs.add(repo_key)
                repo_streams = selected_repo_streams.get(repo_key)
                if repo_streams is None and owner_name.casefold() != org.casefold():
                    repo_key = (org.casefold(), repo_name.casefold())
                    repo_streams = selected_repo_streams.get(repo_key)
                embedded_properties = repository.get("custom_properties")
                if (isinstance(embedded_properties, dict)
                        and (not repo_streams or not any(value.casefold() != "unmapped" for value in repo_streams))):
                    try:
                        embedded_streams = property_value_streams(embedded_properties, property_names)
                    except ValueError:
                        embedded_streams = ("Unmapped",)
                    if any(value.casefold() != "unmapped" for value in embedded_streams):
                        repo_streams = embedded_streams
                if repo_streams is None:
                    LOG.debug("Team %s/%s repository %s/%s is outside the selected repository scope",
                              org, slug, owner_name, repo_name)
                    continue
                matched_repository_refs.add(repo_key)
                if any(value.casefold() != "unmapped" for value in repo_streams):
                    mapped_repository_refs.add(repo_key)
                team_streams.update(repo_streams)

            for member in team_members:
                login = member.get("login")
                if not login or (exclude_bots and is_bot(member)):
                    continue
                teams_by_user[(org, login)].add(name)
                streams_by_user[(org, login.casefold())].update(team_streams)

        for (owner, repo), report in reports.items():
            repo_key = (owner.casefold(), repo.casefold())
            if owner.casefold() != org.casefold() or repo_key in team_repository_refs:
                continue
            if not any(stream.casefold() != "unmapped" for stream in report.value_streams):
                continue
            try:
                repo_teams = client.paginate(owner, f"/repos/{owner}/{repo}/teams")
                for repo_team in repo_teams:
                    slug = repo_team.get("slug")
                    name = repo_team.get("name") or slug
                    if not slug or not name:
                        continue
                    team_key = slug.casefold()
                    if team_key not in team_members_by_slug:
                        try:
                            team_members = list(client.paginate(
                                org, f"/orgs/{org}/teams/{slug}/members"))
                        except (GitHubAPIError, requests.RequestException) as exc:
                            LOG.warning("Could not list members for repository team %s/%s: %s", org, slug, exc)
                            failures.append((f"{org} (team {slug} members)", str(exc)))
                            org_team_complete[org] = False
                            continue
                        team_members_by_slug[team_key] = (
                            name,
                            {member["login"] for member in team_members
                             if member.get("login") and not (exclude_bots and is_bot(member))},
                        )
                    team_name, member_logins = team_members_by_slug[team_key]
                    for login in member_logins:
                        teams_by_user[(org, login)].add(team_name)
                        streams_by_user[(org, login.casefold())].update(report.value_streams)
                    team_repository_refs.add(repo_key)
                    matched_repository_refs.add(repo_key)
                    mapped_repository_refs.add(repo_key)
            except (GitHubAPIError, requests.RequestException) as exc:
                reason = ("GitHub App needs repository Administration: read permission to list repository teams"
                          if isinstance(exc, GitHubAPIError) and exc.status == 403 else str(exc))
                LOG.warning("Could not list repository teams for %s/%s: %s", owner, repo, reason)
                failures.append((f"{org} (repository {repo} team list)", reason))
                org_team_complete[org] = False

        LOG.info("Organization %s: %d visible teams; %d team repositories, %d in selected report, %d with value streams",
             org, len(teams), len(team_repository_refs), len(matched_repository_refs), len(mapped_repository_refs))

    return teams_by_user, streams_by_user, org_team_complete


# --------------------------------------------------------------------------- Actions self-hosted runners

@dataclass
class RunnerInfo:
    id: int
    name: str
    os: str
    status: str
    busy: bool
    labels: Tuple[str, ...]
    runner_group_id: Optional[int] = None
    runner_group_name: str = "Default"
    org: str = ""
    repo: Optional[str] = None
    arch: str = "x64"
    runner_version: str = "Unknown"
    ephemeral: bool = False


@dataclass
class RunnerJobMetrics:
    total_jobs: int = 0
    successful_jobs: int = 0
    failed_jobs: int = 0
    cancelled_jobs: int = 0
    timed_out_jobs: int = 0
    other_jobs: int = 0
    total_duration_seconds: float = 0.0
    total_queue_seconds: float = 0.0
    queue_durations: List[float] = field(default_factory=list)
    job_durations: List[float] = field(default_factory=list)
    first_job_time: Optional[datetime] = None
    last_job_time: Optional[datetime] = None
    repos_served: Set[str] = field(default_factory=set)

    @property
    def completed_jobs(self) -> int:
        return self.successful_jobs + self.failed_jobs + self.cancelled_jobs + self.timed_out_jobs + self.other_jobs

    @property
    def success_rate(self) -> Any:
        return round(self.successful_jobs / self.completed_jobs * 100, 2) if self.completed_jobs else "N/A"

    @property
    def failure_rate(self) -> Any:
        return round((self.failed_jobs + self.timed_out_jobs) / self.completed_jobs * 100, 2) if self.completed_jobs else "N/A"

    @property
    def avg_duration_minutes(self) -> Any:
        return round(self.total_duration_seconds / len(self.job_durations) / 60, 2) if self.job_durations else "N/A"

    @property
    def median_duration_minutes(self) -> Any:
        if not self.job_durations:
            return "N/A"
        sorted_d = sorted(self.job_durations)
        mid = len(sorted_d) // 2
        med = sorted_d[mid] if len(sorted_d) % 2 != 0 else (sorted_d[mid - 1] + sorted_d[mid]) / 2.0
        return round(med / 60.0, 2)

    @property
    def min_duration_minutes(self) -> Any:
        return round(min(self.job_durations) / 60.0, 2) if self.job_durations else "N/A"

    @property
    def max_duration_minutes(self) -> Any:
        return round(max(self.job_durations) / 60.0, 2) if self.job_durations else "N/A"

    @property
    def total_compute_hours(self) -> float:
        return round(self.total_duration_seconds / 3600, 2)

    @property
    def avg_queue_seconds(self) -> Any:
        return round(self.total_queue_seconds / len(self.queue_durations), 2) if self.queue_durations else "N/A"

    @property
    def p95_queue_seconds(self) -> Any:
        if not self.queue_durations:
            return "N/A"
        sorted_q = sorted(self.queue_durations)
        idx = max(0, min(len(sorted_q) - 1, math.ceil(0.95 * len(sorted_q)) - 1))
        return round(sorted_q[idx], 2)

    @property
    def queue_sla_breaches(self) -> int:
        return sum(1 for q in self.queue_durations if q > 60.0)

    @property
    def queue_sla_breach_rate(self) -> Any:
        return round(self.queue_sla_breaches / len(self.queue_durations) * 100, 2) if self.queue_durations else "N/A"


def collect_runners_and_metrics(
    client: GitHubClient,
    orgs: List[str],
    targets: List[Tuple[str, str]],
    since: datetime,
    end: datetime,
    cfg: Config,
    failures: List[Tuple[str, str]]
) -> Tuple[List[RunnerInfo], Dict[Tuple[str, str], RunnerJobMetrics], Dict[str, RunnerJobMetrics]]:
    runners: List[RunnerInfo] = []
    runner_metrics: Dict[Tuple[str, str], RunnerJobMetrics] = defaultdict(RunnerJobMetrics)
    org_runner_metrics: Dict[str, RunnerJobMetrics] = defaultdict(RunnerJobMetrics)

    known_runner_names: Set[str] = set()
    known_runner_ids: Set[int] = set()

    for org in orgs:
        groups: Dict[int, str] = {}
        try:
            for item in client.paginate_enveloped(org, f"/orgs/{org}/actions/runner-groups", "runner_groups"):
                if isinstance(item, dict) and "id" in item:
                    groups[item["id"]] = item.get("name", "Default")
        except (GitHubAPIError, requests.RequestException) as exc:
            LOG.debug("Could not retrieve runner groups for %s: %s", org, exc)

        try:
            for item in client.paginate_enveloped(org, f"/orgs/{org}/actions/runners", "runners"):
                if not isinstance(item, dict) or "id" not in item:
                    continue
                r_id = item["id"]
                name = item.get("name", f"runner-{r_id}")
                os_name = item.get("os", "Unknown")
                status = str(item.get("status", "unknown")).capitalize()
                busy = bool(item.get("busy", False))
                raw_labels = item.get("labels", [])
                labels = tuple(sorted(l.get("name") for l in raw_labels if isinstance(l, dict) and l.get("name")))
                g_id = item.get("runner_group_id")
                g_name = groups.get(g_id, "Default") if g_id else "Default"
                arch = item.get("arch", "x64")
                version = item.get("runner_version") or item.get("version") or "Unknown"
                ephemeral = bool(item.get("ephemeral", False))
                runners.append(RunnerInfo(
                    id=r_id, name=name, os=os_name, status=status, busy=busy,
                    labels=labels, runner_group_id=g_id, runner_group_name=g_name,
                    org=org, repo=None, arch=arch, runner_version=version, ephemeral=ephemeral
                ))
                known_runner_names.add(name.casefold())
                known_runner_ids.add(r_id)
            LOG.info("Organization %s: %d self-hosted runner(s) discovered",
                     org, len([r for r in runners if r.org.casefold() == org.casefold() and not r.repo]))
        except (GitHubAPIError, requests.RequestException) as exc:
            reason = str(exc)
            if isinstance(exc, GitHubAPIError) and exc.status == 403:
                reason = "GitHub App requires 'Organization self-hosted runners: Read-only' permission"
                LOG.warning("Self-hosted runners not accessible for %s: %s", org, reason)
                failures.append((f"{org} (Actions self-hosted runners)", reason))
            elif isinstance(exc, GitHubAPIError) and exc.status == 404:
                LOG.info("Self-hosted runners not configured for org %s: %s", org, reason)
            else:
                failures.append((f"{org} (Actions self-hosted runners)", reason))

    for org, repo in targets:
        try:
            for item in client.paginate_enveloped(org, f"/repos/{org}/{repo}/actions/runners", "runners"):
                if not isinstance(item, dict) or "id" not in item:
                    continue
                r_id = item["id"]
                name = item.get("name", f"runner-{r_id}")
                os_name = item.get("os", "Unknown")
                status = str(item.get("status", "unknown")).capitalize()
                busy = bool(item.get("busy", False))
                raw_labels = item.get("labels", [])
                labels = tuple(sorted(l.get("name") for l in raw_labels if isinstance(l, dict) and l.get("name")))
                arch = item.get("arch", "x64")
                version = item.get("runner_version") or item.get("version") or "Unknown"
                ephemeral = bool(item.get("ephemeral", False))
                runners.append(RunnerInfo(
                    id=r_id, name=name, os=os_name, status=status, busy=busy,
                    labels=labels, runner_group_id=None, runner_group_name="Repository-level",
                    org=org, repo=repo, arch=arch, runner_version=version, ephemeral=ephemeral
                ))
                known_runner_names.add(name.casefold())
                known_runner_ids.add(r_id)
        except (GitHubAPIError, requests.RequestException):
            pass

    for org, repo in targets:
        try:
            runs_checked = 0
            for run_item in client.paginate_enveloped(org, f"/repos/{org}/{repo}/actions/runs", "workflow_runs", params={"status": "completed"}):
                created_str = run_item.get("created_at") or run_item.get("run_started_at")
                if not created_str:
                    continue
                run_time = parse_ts(created_str)
                if run_time < since:
                    break
                if run_time > end:
                    continue
                runs_checked += 1
                if runs_checked > cfg.max_runner_runs_per_repo:
                    break

                run_id = run_item["id"]
                try:
                    jobs = list(client.paginate_enveloped(org, f"/repos/{org}/{repo}/actions/runs/{run_id}/jobs", "jobs"))
                except (GitHubAPIError, requests.RequestException):
                    continue

                for job in jobs:
                    if not isinstance(job, dict):
                        continue
                    job_labels = [str(lbl).casefold() for lbl in job.get("labels", []) if isinstance(lbl, str)]
                    r_id = job.get("runner_id")
                    r_name = job.get("runner_name") or ""

                    is_self_hosted = (
                        "self-hosted" in job_labels
                        or (r_id and r_id in known_runner_ids)
                        or (r_name and r_name.casefold() in known_runner_names)
                    )
                    if not is_self_hosted:
                        continue

                    matched_runner_name = r_name or (f"runner-{r_id}" if r_id else "self-hosted-runner")
                    if matched_runner_name.casefold() not in known_runner_names:
                        known_runner_names.add(matched_runner_name.casefold())
                        runners.append(RunnerInfo(
                            id=r_id or 0, name=matched_runner_name, os=job.get("os", "Linux"),
                            status="Ephemeral", busy=False, labels=tuple(job.get("labels", ["self-hosted"])),
                            runner_group_name=job.get("runner_group_name") or "Default",
                            org=org, repo=None, arch=job.get("arch", "x64"), runner_version="Unknown", ephemeral=True
                        ))

                    created_at = parse_ts(job["created_at"]) if job.get("created_at") else None
                    started_at = parse_ts(job["started_at"]) if job.get("started_at") else None
                    completed_at = parse_ts(job["completed_at"]) if job.get("completed_at") else None
                    conclusion = str(job.get("conclusion", "other")).casefold()

                    queue_sec = max(0.0, (started_at - created_at).total_seconds()) if (started_at and created_at) else None
                    duration_sec = max(0.0, (completed_at - started_at).total_seconds()) if (completed_at and started_at) else None
                    job_end = completed_at or started_at

                    for metrics_bucket in (runner_metrics[(org, matched_runner_name)], org_runner_metrics[org]):
                        metrics_bucket.total_jobs += 1
                        metrics_bucket.repos_served.add(repo)
                        if conclusion == "success":
                            metrics_bucket.successful_jobs += 1
                        elif conclusion == "failure":
                            metrics_bucket.failed_jobs += 1
                        elif conclusion == "timed_out":
                            metrics_bucket.timed_out_jobs += 1
                        elif conclusion == "cancelled":
                            metrics_bucket.cancelled_jobs += 1
                        else:
                            metrics_bucket.other_jobs += 1

                        if duration_sec is not None:
                            metrics_bucket.total_duration_seconds += duration_sec
                            metrics_bucket.job_durations.append(duration_sec)
                        if queue_sec is not None:
                            metrics_bucket.total_queue_seconds += queue_sec
                            metrics_bucket.queue_durations.append(queue_sec)
                        if started_at:
                            metrics_bucket.first_job_time = (
                                min(metrics_bucket.first_job_time, started_at)
                                if metrics_bucket.first_job_time else started_at
                            )
                        if job_end:
                            metrics_bucket.last_job_time = (
                                max(metrics_bucket.last_job_time, job_end)
                                if metrics_bucket.last_job_time else job_end
                            )
        except (GitHubAPIError, requests.RequestException) as exc:
            reason = str(exc)
            if not (isinstance(exc, GitHubAPIError) and exc.status in (403, 404)):
                failures.append((f"{org}/{repo} (Actions jobs)", reason))

    return runners, runner_metrics, org_runner_metrics


def build_runner_rows(
    runners: List[RunnerInfo],
    runner_metrics: Dict[Tuple[str, str], RunnerJobMetrics],
    org_runner_metrics: Dict[str, RunnerJobMetrics],
    orgs: Optional[List[str]] = None
) -> Tuple[List[list], List[list]]:
    runner_rows: List[list] = []
    org_runner_rows: List[list] = []

    for r in sorted(runners, key=lambda x: (x.org.casefold(), x.name.casefold())):
        metrics = runner_metrics.get((r.org, r.name), RunnerJobMetrics())
        scope = f"Repo: {r.repo}" if r.repo else "Organization"
        runner_rows.append([
            r.org,
            r.id,
            r.name,
            scope,
            r.runner_group_name,
            r.os,
            r.arch,
            r.runner_version,
            "Ephemeral (Autoscaled)" if r.ephemeral else "Persistent (Static VM)",
            r.status,
            "Busy" if r.busy else "Idle",
            "; ".join(r.labels) if r.labels else "None",
            len(metrics.repos_served),
            metrics.total_jobs,
            metrics.successful_jobs,
            metrics.failed_jobs,
            metrics.cancelled_jobs,
            metrics.timed_out_jobs,
            metrics.success_rate,
            metrics.total_compute_hours,
            metrics.avg_duration_minutes,
            metrics.median_duration_minutes,
            metrics.min_duration_minutes,
            metrics.max_duration_minutes,
            metrics.avg_queue_seconds,
            metrics.p95_queue_seconds,
            metrics.queue_sla_breaches,
            fmt_timestamp(metrics.first_job_time) if metrics.first_job_time else "No activity",
            fmt_timestamp(metrics.last_job_time) if metrics.last_job_time else "No activity",
        ])

    all_orgs = sorted(set(orgs or []) | {r.org for r in runners} | set(org_runner_metrics), key=str.casefold)
    for org in all_orgs:
        org_runners = [r for r in runners if r.org.casefold() == org.casefold()]
        total_count = len(org_runners)
        org_level_count = sum(1 for r in org_runners if not r.repo)
        repo_level_count = sum(1 for r in org_runners if r.repo)
        online_count = sum(1 for r in org_runners if r.status.casefold() == "online")
        offline_count = sum(1 for r in org_runners if r.status.casefold() in ("offline", "ephemeral"))
        busy_count = sum(1 for r in org_runners if r.busy)
        ephemeral_count = sum(1 for r in org_runners if r.ephemeral)
        availability = round(online_count / total_count * 100, 2) if total_count else "N/A"
        metrics = org_runner_metrics.get(org, RunnerJobMetrics())

        org_runner_rows.append([
            org,
            total_count,
            org_level_count,
            repo_level_count,
            online_count,
            offline_count,
            availability,
            busy_count,
            ephemeral_count,
            len(metrics.repos_served),
            metrics.total_jobs,
            metrics.successful_jobs,
            metrics.failed_jobs,
            metrics.cancelled_jobs,
            metrics.timed_out_jobs,
            metrics.success_rate,
            metrics.total_compute_hours,
            metrics.avg_queue_seconds,
            metrics.p95_queue_seconds,
            metrics.queue_sla_breaches,
            metrics.queue_sla_breach_rate,
        ])

    return runner_rows, org_runner_rows


# --------------------------------------------------------------------------- report building

def classify(active_days: int, cfg: Config) -> str:
    if active_days >= cfg.active_min_days:
        return "Active"
    if active_days >= cfg.moderate_min_days:
        return "Moderate"
    return "Low" if active_days else "Inactive"


def build_rows(results: Dict[Tuple[str, str], Dict[str, Activity]], members: Dict[str, Set[str]],
               cfg: Config, end: datetime, user_streams: Dict[Tuple[str, str], str],
               teams_by_user: Dict[Tuple[str, str], Set[str]],
               org_team_complete: Dict[str, bool]) -> Tuple[List[list], List[list]]:
    per_user: Dict[Tuple[str, str], Activity] = defaultdict(Activity)
    repos_by_user: Dict[Tuple[str, str], Set[str]] = defaultdict(set)
    teams_by_user_normalized: Dict[Tuple[str, str], Set[str]] = defaultdict(set)
    repo_rows: List[list] = []

    for (org, login), teams in teams_by_user.items():
        teams_by_user_normalized[(org, login.casefold())].update(teams)

    for (org, repo), users in sorted(results.items(), key=lambda item: (item[0][0].lower(), item[0][1].lower())):
        for login, act in sorted(users.items(), key=lambda item: (-item[1].total, item[0].lower())):
            per_user[(org, login)].merge(act)
            repos_by_user[(org, login)].add(repo)
            repo_rows.append([org, repo, login, act.total, *(act.counts[m] for m in METRIC_KEYS),
                              len(act.days), fmt_timestamp(act.last)])

    if cfg.include_inactive_members:
        for org, logins in members.items():
            for login in logins:
                per_user.setdefault((org, login), Activity())
    for org, login in teams_by_user:
        per_user.setdefault((org, login), Activity())

    user_rows: List[list] = []
    for (org, login), act in per_user.items():
        member = ("Yes" if login in members[org] else "No") if org in members else "Unknown"
        teams = sorted(teams_by_user_normalized.get((org, login.casefold()), ()), key=str.casefold)
        team_value = "; ".join(teams)
        if not org_team_complete.get(org, False):
            team_value = f"{team_value}; Unknown" if team_value else "Unknown"
        elif not team_value:
            team_value = "None"
        active_months = {(active_day.year, active_day.month) for active_day in act.days}
        avg_commits_per_month = round(act.counts["commits"] / len(active_months), 2) if active_months else None
        user_rows.append([
            org, login, classify(len(act.days), cfg), len(act.days), act.total,
            *(act.counts[m] for m in METRIC_KEYS), avg_commits_per_month,
            len(repos_by_user.get((org, login), ())), fmt_timestamp(act.first), fmt_timestamp(act.last),
            (end - act.last).days if act.last else None, member, team_value,
            user_streams.get((org, login.casefold()), ""),
        ])
    user_rows.sort(key=lambda r: (r[0].lower(), STATUS_ORDER[r[STATUS_IDX]], -r[TOTAL_IDX], r[1].lower()))
    return user_rows, repo_rows


def build_user_streams(reports: Dict[Tuple[str, str], RepositoryReport], since: datetime,
                       end: datetime,
                       team_streams_by_user: Optional[Dict[Tuple[str, str], Set[str]]] = None) -> Dict[Tuple[str, str], str]:
    streams_by_user: Dict[Tuple[str, str], Set[str]] = defaultdict(set)
    for key, streams in (team_streams_by_user or {}).items():
        streams_by_user[key].update(streams)
    for (org, repo), report in reports.items():
        for login, _, when in report.events:
            if login and since <= when <= end:
                streams_by_user[(org, login.casefold())].update(report.value_streams)
    user_streams = {}
    for key, streams in streams_by_user.items():
        mapped_streams = {stream for stream in streams if stream.casefold() != "unmapped"}
        user_streams[key] = "; ".join(sorted(mapped_streams or streams, key=str.casefold))
    return user_streams


def build_org_rows(user_rows: List[list], results: Dict[Tuple[str, str], Dict[str, Activity]]) -> List[list]:
    repo_counts = Counter(org for org, _ in results)
    stats: Dict[str, Counter] = defaultdict(Counter)
    for row in user_rows:
        org_stats = stats[row[0]]
        org_stats[row[STATUS_IDX]] += 1
        org_stats["total"] += row[TOTAL_IDX]
        org_stats["contributors"] += 1 if row[TOTAL_IDX] else 0
    orgs = sorted(set(repo_counts) | set(stats), key=str.lower)
    return [[org, repo_counts[org], stats[org]["contributors"], *(stats[org][s] for s in STATUSES),
             stats[org]["total"]] for org in orgs]


def history_scope(cfg: Config, entries: List[RepoEntry]) -> str:
    scope = [cfg.api_url, cfg.include_archived, cfg.include_forks, cfg.exclude_bots,
             sorted((entry.org.lower(), entry.repo.lower(), entry.enabled) for entry in entries)]
    return hashlib.sha256(json.dumps(scope, sort_keys=True).encode("utf-8")).hexdigest()


def load_history(path: Path) -> Dict[str, Any]:
    if not path.exists():
        return {"version": 1, "snapshots": []}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, dict) or data.get("version") != 1 or not isinstance(data.get("snapshots"), list):
            raise ValueError("unsupported history format")
        for snapshot in data["snapshots"]:
            parse_ts(snapshot["at"])
            if not isinstance(snapshot["scope"], str) or not isinstance(snapshot["organizations"], dict):
                raise ValueError("invalid snapshot")
            for record in snapshot["organizations"].values():
                if (not isinstance(record["members"], list)
                        or any(not isinstance(member, str) or not re.fullmatch(r"[0-9a-f]{64}", member)
                               for member in record["members"])
                        or isinstance(record["repositories"], bool) or not isinstance(record["repositories"], int)
                        or record["repositories"] < 0):
                    raise ValueError("invalid snapshot counts")
        return data
    except (ValueError, KeyError, TypeError, AttributeError) as exc:
        raise ConfigError("Invalid adoption history; restore a valid state file or configure a new history_file") from exc


def membership_snapshot(scope: str, end: datetime, members: Dict[str, Set[str]],
                        reports: Dict[Tuple[str, str], RepositoryReport]) -> Dict[str, Any]:
    repo_counts = Counter(org.lower() for org, _ in reports)
    return {"scope": scope, "at": end.isoformat(), "organizations": {
        org.lower(): {"members": sorted(hashlib.sha256(login.lower().encode("utf-8")).hexdigest() for login in logins),
                      "repositories": repo_counts[org.lower()]}
        for org, logins in members.items()}}


def save_history(path: Path, data: Dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    try:
        with temporary.open("x", encoding="utf-8") as handle:
            json.dump(data, handle, indent=2)
            handle.flush()
            os.fsync(handle.fileno())
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def period_start(value: datetime, grain: str) -> datetime:
    local = value.astimezone(IST).replace(hour=0, minute=0, second=0, microsecond=0)
    return local - timedelta(days=local.weekday()) if grain == "Weekly" else local.replace(day=1)


def next_period(value: datetime, grain: str) -> datetime:
    if grain == "Weekly":
        return value + timedelta(days=7)
    return value.replace(year=value.year + 1, month=1) if value.month == 12 else value.replace(month=value.month + 1)


def percentage_change(current: int, previous: Optional[int]) -> Any:
    return round((current - previous) / previous * 100, 2) if previous else "N/A"


def snapshot_totals(snapshot: Dict[str, Any], org: Optional[str] = None) -> Tuple[Set[str], int]:
    records = snapshot["organizations"]
    selected = [records[org.lower()]] if org and org.lower() in records else ([] if org else list(records.values()))
    users = set().union(*(set(record["members"]) for record in selected))
    return users, sum(record["repositories"] for record in selected)


def build_management_tables(reports: Dict[Tuple[str, str], RepositoryReport], members: Dict[str, Set[str]],
                            since: datetime, end: datetime, snapshots: List[Dict[str, Any]],
                            organization_status: Dict[str, bool],
                            audit_joined_users: Optional[Dict[str, Set[str]]] = None,
                            runners: Optional[List[RunnerInfo]] = None,
                            org_runner_metrics: Optional[Dict[str, RunnerJobMetrics]] = None
                            ) -> Dict[str, Tuple[Tuple[str, ...], List[list]]]:
    tables: Dict[str, Tuple[Tuple[str, ...], List[list]]] = {}
    kpis: List[list] = []
    inventory: List[list] = []
    created_dates = {key: parse_ts(result.metadata["created_at"]) for key, result in reports.items()}
    dated_snapshots = sorted((parse_ts(snapshot["at"]), index, snapshot) for index, snapshot in enumerate(snapshots))
    baseline = next((snapshot for when, _, snapshot in reversed(dated_snapshots) if when <= since), None)
    streams: Dict[str, Dict[str, Any]] = defaultdict(lambda: {
        "repos": 0, "members": set(), "contributors": set(), "private": 0, "internal": 0, "public": 0,
        "organizations": set(), "repository_refs": set(), "member_coverage": True})
    for (org, repo), result in reports.items():
        stream = "; ".join(result.value_streams)
        visibility = result.metadata.get("visibility", "Unknown")
        events = [event for event in result.events if since <= event[2] <= end]
        counts = Counter(metric for _, metric, _ in events)
        contributors = {login for login, _, _ in events if login}
        org_members = {login.lower() for login in members.get(org, set())}
        for value_stream in result.value_streams:
            stats = streams[value_stream]
            stats["repos"] += 1
            stats["organizations"].add(org)
            stats["repository_refs"].add(f"{org}/{repo}")
            if visibility in {"private", "internal", "public"}:
                stats[visibility] += 1
            stats["contributors"].update(contributors)
            stats["members"].update(contributors & org_members)
            stats["member_coverage"] = stats["member_coverage"] and org in members
        activity_available = result.activity_complete
        inventory.append([
            org, repo, stream, visibility, fmt_timestamp(created_dates[(org, repo)]),
            "Yes" if since <= created_dates[(org, repo)] <= end else "No",
            "Yes" if result.metadata.get("archived") else "No", result.metadata.get("default_branch", ""),
            "; ".join(result.protected_branches) if result.protected == "Yes" else
            "None" if result.protected == "No" else "Unknown",
            result.codeowners, result.codeowners_path,
            ("Yes" if any(when >= end - timedelta(days=7) for _, _, when in result.events) else "No")
            if activity_available else "Unknown",
            len(events) if activity_available else "N/A", len(contributors) if activity_available else "N/A",
            *(counts[metric] if activity_available else "N/A" for metric in METRIC_KEYS),
            round(result.cohort_merged / result.cohort_opened * 100, 2) if result.cohort_opened else "N/A",
            round(sum(result.review_hours) / len(result.review_hours), 2) if result.review_hours else "N/A",
            len(result.review_hours), "Complete" if activity_available and not result.errors else "Partial",
        ])
    inventory.sort(key=lambda row: (-(row[12] if isinstance(row[12], int) else -1), row[0].lower(), row[1].lower()))
    tables["Repository Inventory"] = ((
        "Organization", "Repository", "Value Stream", "Visibility", "Created (IST)", "New in Period",
        "Archived", "Default Branch", "Protected Branches", "CODEOWNERS Present", "CODEOWNERS Path",
        "Activity Last 7 Days", "Total Activities", "Active Contributors", *METRIC_LABELS,
        "PR Cohort Merge Rate (%)", "Avg First Review Hours", "First Review Sample PRs", "Coverage"), inventory)
    inventory_headers = tables["Repository Inventory"][0]
    top_headers = (
        "Organization", "Rank", "Repository", "Value Stream", "Total Activities", "Active Contributors",
        *METRIC_LABELS, "Activity Last 7 Days", "Coverage",
    )
    top_rows = []
    for org in sorted({row[0] for row in inventory}, key=str.casefold):
        org_repositories = [row for row in inventory if row[0] == org]
        org_repositories.sort(key=lambda row: (
            -(row[inventory_headers.index("Total Activities")]
              if isinstance(row[inventory_headers.index("Total Activities")], int) else -1),
            row[1].casefold()))
        for rank, row in enumerate(org_repositories[:10], start=1):
            top_rows.append([org, rank, row[inventory_headers.index("Repository")],
                             row[inventory_headers.index("Value Stream")],
                             row[inventory_headers.index("Total Activities")],
                             row[inventory_headers.index("Active Contributors")],
                             *(row[inventory_headers.index(label)] for label in METRIC_LABELS),
                             row[inventory_headers.index("Activity Last 7 Days")],
                             row[inventory_headers.index("Coverage")]])
    tables["Top Active Repositories"] = (top_headers, top_rows)
    tables["Value Streams"] = (("Value Stream", "Organization Names", "Repository Names (Org/Repo)",
                                "Repositories", "Active Org Members", "Active Contributors",
                                "Private", "Internal", "Public"), [
        [name, "; ".join(sorted(data["organizations"], key=str.casefold)),
         "; ".join(sorted(data["repository_refs"], key=str.casefold)), data["repos"],
         len(data["members"]) if data["member_coverage"] else "N/A",
         len(data["contributors"]),
         data["private"], data["internal"], data["public"]] for name, data in sorted(streams.items())])

    organizations = sorted({org for org, _ in reports} | set(members) | set(organization_status), key=str.lower)
    for org in organizations:
        selected = [result for (owner, _), result in reports.items() if owner == org]
        events = [event for result in selected for event in result.events if since <= event[2] <= end]
        counts = Counter(metric for _, metric, _ in events)
        contributors = {login for login, _, _ in events if login}
        org_members = {login.lower() for login in members.get(org, set())}
        member_coverage = org in members
        rolling = {login for result in selected for login, _, when in result.events
                   if login and end - timedelta(days=30) <= when <= end}
        month_start = period_start(end, "Monthly")
        monthly = {login for result in selected for login, _, when in result.events if login and month_start <= when <= end}
        hours = [hours for result in selected for hours in result.review_hours]
        opened = sum(result.cohort_opened for result in selected)
        merged = sum(result.cohort_merged for result in selected)
        def add(category: str, metric: str, value: Any, definition: str) -> None:
            kpis.append([org, category, metric, value, definition])
        add("User Adoption", "Current onboarded org members", len(org_members) if member_coverage else "N/A",
            "How many people are currently listed as members of this organization. This is not a paid-seat or sign-in count.")
        add("User Adoption", "Active onboarded members", len(org_members & contributors) if member_coverage else "N/A",
            "Current organization members who made at least one tracked contribution during this report period.")
        add("User Adoption", "Monthly active users (month to date)", len(monthly),
            "Distinct people with tracked activity so far this calendar month (IST), including people who are not members.")
        add("User Adoption", "Rolling 30-day active users", len(rolling),
            "Distinct people with tracked activity during the last 30 days, including non-members.")
        onboarding_users = (audit_joined_users or {}).get(org.casefold(), set())
        onboarding_value: Any = len(onboarding_users) if audit_joined_users is not None else "Unavailable"
        onboarding_definition = (
            "Distinct users with org.add_member audit events during the report period. The CSV must cover the full period. "
            "A removed and re-added user can count; this does not establish first-ever membership."
            if audit_joined_users is not None else
            "Unavailable without an organization audit-log CSV covering the full report period. The current member list has no join dates."
        )
        add("User Adoption", "New users onboarded (exact)", onboarding_value, onboarding_definition)
        observed: Any = "N/A"
        if baseline and member_coverage:
            known, _ = snapshot_totals(baseline, org)
            additions: Set[str] = set()
            for when, _, snapshot in dated_snapshots:
                if since < when <= end:
                    current, _ = snapshot_totals(snapshot, org)
                    additions.update(current - known)
                    known |= current
            observed = len(additions)
        add("User Adoption", "Newly observed org members", observed,
            "Members first seen in a saved snapshot during this period. This is an observed change, not the person's actual join date; an earlier snapshot is needed.")
        add("Repository Landscape", "Repositories hosted (selected scope)", len(selected),
            "Repositories selected in repos.csv that this App can access. This may not include every repository in the enterprise.")
        add("Repository Landscape", "New repositories created", sum(since <= created <= end
            for (owner, _), created in created_dates.items() if owner == org),
            "Selected repositories created during this period that still exist here. Deleted or transferred repositories are not counted.")
        for visibility in ("private", "internal", "public"):
            add("Repository Landscape", f"{visibility.title()} repositories",
                sum(item.metadata.get("visibility") == visibility for item in selected),
                f"Selected repositories currently set to {visibility} visibility.")
        add("Developer Activity", "Repositories active in last 7 days", sum(any(when >= end - timedelta(days=7)
            for _, _, when in item.events) for item in selected),
            "Repositories with at least one tracked contribution in the last seven days. This does not measure logins or every kind of push.")
        add("Developer Activity", "Total commit volume", counts["commits"],
            "Commit records on default branches during the report period, counted by author date. Some commits may not be linked to a GitHub user; the bot filter applies.")
        add("Developer Activity", "Active contributors", len(contributors),
            "Distinct GitHub users with at least one tracked contribution, including people outside the organization.")
        add("Developer Activity", "Average commits per active user",
            round(counts["commits"] / len(contributors), 2) if contributors else "N/A",
            "Default-branch commits divided by distinct active contributors. Unlinked commits can be included in the commit count but not the user count.")
        add("Developer Activity", "Top Active Repositories by Activity", "See Top Active Repositories sheet",
            "The sheet lists this organization's ten repositories with the most tracked events in the report period. It is an activity count, not a productivity score.")
        add("Collaboration", "Pull requests created", counts["prs_opened"], "Pull requests opened during the report period.")
        add("Collaboration", "Pull requests merged", counts["prs_merged"],
            "Pull requests merged during the report period, even if they were opened earlier.")
        add("Collaboration", "PR merge rate (%)", round(merged / opened * 100, 2) if opened else "N/A",
            "Of the PRs opened during this period, the percentage merged by the report cutoff. Shows N/A if no PRs were opened.")
        add("Collaboration", "Code reviews completed", counts["reviews"],
            "Submitted pull-request reviews during the period. Draft reviews are excluded; submitted reviews count even if later dismissed.")
        add("Collaboration", "Average PR review turnaround (hours)", round(sum(hours) / len(hours), 2) if hours else "N/A",
            "Average time from opening a PR to its first submitted review by someone else. Unreviewed PRs are excluded; weekends and draft time count.")
        add("Collaboration", "First review sample PRs", len(hours),
            "Number of PRs used to calculate the average first-review time.")
        for attribute, title in (("protected", "Repositories with branch protection"), ("codeowners", "Repositories with CODEOWNERS")):
            add("Security & Governance", title, sum(getattr(item, attribute) == "Yes" for item in selected),
                "Repositories with at least one protected branch." if attribute == "protected" else
                "Repositories with a CODEOWNERS file in a recognized location on the default branch. The file's rules are not validated.")
            add("Security & Governance", f"{title} - unknown", sum(getattr(item, attribute) == "Unknown" for item in selected),
                "Repositories where GitHub did not provide enough information to confirm this check. They are not counted as having no protection/file.")
        add("Coverage", "Activity scans completed", sum(item.activity_complete for item in selected),
            "Number of this organization's selected repositories whose activity scan completed successfully.")
        add("Coverage", "Run status", "Complete" if organization_status.get(org, False) else "Partial",
            "Complete means this organization's selected repositories and member list were collected without errors. A failure in another organization does not change this status.")
        org_runners = [r for r in (runners or []) if r.org.casefold() == org.casefold()]
        r_metrics = (org_runner_metrics or {}).get(org, RunnerJobMetrics())
        total_r = len(org_runners)
        online_r = sum(1 for r in org_runners if str(r.status).casefold() == "online")
        ephemeral_r = sum(1 for r in org_runners if r.ephemeral)
        availability = round(online_r / total_r * 100, 2) if total_r else "N/A"
        ephemeral_pct = round(ephemeral_r / total_r * 100, 2) if total_r else "N/A"
        add("Actions Self-Hosted Runners", "Total self-hosted runners registered", total_r,
            "Total number of self-hosted runners configured across the organization and its scanned repositories.")
        add("Actions Self-Hosted Runners", "Online runner availability (%)", availability,
            "Percentage of registered self-hosted runners currently reporting online and ready to accept jobs.")
        add("Actions Self-Hosted Runners", "Ephemeral autoscaled runners (%)", ephemeral_pct,
            "Percentage of runners operating in ephemeral / auto-scaling mode (e.g. Actions Runner Controller / K8s).")
        add("Actions Self-Hosted Runners", "Self-hosted jobs executed", r_metrics.total_jobs,
            "Number of completed workflow jobs executed on self-hosted runners during the lookback period.")
        add("Actions Self-Hosted Runners", "Total self-hosted compute time (hours)", r_metrics.total_compute_hours,
            "Total machine compute hours executed across all self-hosted runners.")
        add("Actions Self-Hosted Runners", "Average job queue latency (seconds)", r_metrics.avg_queue_seconds,
            "Mean duration jobs waited in queue before a self-hosted runner picked them up. High values signal capacity saturation.")
        add("Actions Self-Hosted Runners", "P95 job queue latency (seconds)", r_metrics.p95_queue_seconds,
            "95th percentile queue latency. Standard DevOps SLA benchmark measuring peak runner scheduling delay.")
        add("Actions Self-Hosted Runners", "Queue SLA breach rate (% >60s)", r_metrics.queue_sla_breach_rate,
            "Percentage of workflow jobs delayed more than 60 seconds in queue waiting for a self-hosted runner.")
        add("Actions Self-Hosted Runners", "Self-hosted job success rate (%)", r_metrics.success_rate,
            "Percentage of self-hosted workflow jobs completing with success conclusion.")
        add("Actions Self-Hosted Runners", "Repositories utilizing self-hosted runners", len(r_metrics.repos_served),
            "Number of distinct repositories that executed workflow jobs on self-hosted runners.")
        add("Actions Self-Hosted Runners", "Total successful jobs", r_metrics.successful_jobs,
            "Count of completed workflow jobs on self-hosted runners that succeeded.")
        add("Actions Self-Hosted Runners", "Total failed jobs", r_metrics.failed_jobs,
            "Count of workflow jobs on self-hosted runners that failed.")
        add("Actions Self-Hosted Runners", "Total cancelled jobs", r_metrics.cancelled_jobs,
            "Count of workflow jobs on self-hosted runners that were cancelled by users.")
        add("Actions Self-Hosted Runners", "Self-hosted timed out jobs", r_metrics.timed_out_jobs,
            "Number of workflow jobs aborted due to execution timeout on self-hosted runners.")
    tables["Management KPIs"] = (("Scope", "Category", "Metric", "Value", "What this means"), kpis)

    trends: List[list] = []
    for grain in ("Weekly", "Monthly"):
        snapshot_buckets: Dict[datetime, Dict[str, Any]] = {}
        for when, _, snapshot in dated_snapshots:
            if since <= when <= end:
                snapshot_buckets[period_start(when, grain)] = snapshot
        start = period_start(since, grain)
        previous_users: Optional[int] = None
        while start <= end:
            boundary = next_period(start, grain)
            bucket_end = min(boundary, end)
            latest = snapshot_buckets.get(start)
            users, _ = snapshot_totals(latest) if latest else (None, None)
            user_count = len(users) if users is not None else None
            growth = "No growth"
            if user_count is not None and previous_users is not None:
                if previous_users == 0:
                    growth = "No growth" if user_count == 0 else "No baseline"
                else:
                    growth = percentage_change(user_count, previous_users)
            trends.append([
                grain, start.strftime(DATE_FMT), (bucket_end - timedelta(microseconds=1) if boundary <= end else end).strftime(DATE_FMT),
                "Partial" if start < since or boundary > end else "Full",
                growth,
            ])
            previous_users = user_count
            start = boundary
    tables["Growth Trends"] = (("Granularity", "Period Start (IST)", "Period End (IST)", "Period Coverage",
        "Observed User Growth (%)"), trends)
    tables["Metric Notes"] = (("Topic", "What this means"), [
        ["Scope", "This report includes only enabled organizations and repositories in repos.csv that the GitHub App can read. Filters in config.json also apply."],
        ["Adoption", "The member count shows who belongs to the organization now. It does not show paid seats, sign-ins, or when someone joined."],
        ["Onboarding", "New users onboarded (exact) counts distinct users from org.add_member audit events. Supply a complete organization audit-log CSV via input.audit_log_csv; invitations are not joins, and re-added users are not guaranteed to be first-time members."],
        ["History", "A complete run saves a snapshot for later comparisons. Keep adoption_history.json between runs. Without an earlier snapshot, growth is shown as N/A."],
        ["Value streams", "Values come from repository custom properties named VALUE_STREAM or value_stream. Unmapped means no value is set; Unknown means GitHub could not return it. Repositories with multiple values appear in each stream."],
        ["Value-stream users", "Streams come from repositories where a person contributed or from repositories their teams can access. Team membership and team access must be visible to the App."],
        ["Activity", "The report counts default-branch commits, pull requests, reviews, issues, and comments. It does not measure sign-ins, browsing, cloning, or pushes to other branches."],
        ["Totals", "People are matched across organizations by GitHub login. Commits without a linked login count toward commit totals, but cannot be assigned to a person."],
        ["Review turnaround", "This is the time from PR creation to the first submitted review by someone else. Weekends and draft time count; unreviewed PRs are left out. The sample size is shown."],
        ["Merge rate", "This compares PRs opened during the report period with those merged by the report date. The merged total can also include PRs opened earlier."],
        ["Repository growth", "Creation counts include selected repositories that still exist. Deleted or transferred repositories are not available for this comparison. Snapshot changes can include removals."],
        ["Governance", "Protected Branches lists the branches GitHub reports as protected. CODEOWNERS checks only whether a file exists in a recognized default-branch location; it does not validate its rules."],
        ["Ranking", "Repository Inventory is sorted by total activity, highest first. This is a count of events, not a measure of productivity or quality."],
        ["Partial runs", "A failed check can leave totals incomplete. Review Items Not Processed and Coverage before relying on the numbers."],
    ])
    notes = [["All Sheets", topic, definition] for topic, definition in tables["Metric Notes"][1]]
    notes.extend([
        ["Summary", "Report Title", "Fixed report title."],
        ["Summary", "Reporting Period", "Lookback start and report end, formatted in IST."],
        ["Summary", "Generated", "Workbook generation timestamp in IST."],
        ["Summary", "Metric / Value", "Summary KPI name and calculated value; see the corresponding Management KPIs definition."],
        ["Summary", "Organization Breakdown: Organization", "Organization login."],
        ["Summary", "Organization Breakdown: Repositories Scanned", "Repositories with successfully collected activity results."],
        ["Summary", "Organization Breakdown: Contributors", "Distinct users with tracked activity during the reporting period."],
        *[["Summary", f"Organization Breakdown: {status}", f"Users classified {status} by distinct active IST dates using configured thresholds."] for status in STATUSES],
        ["Summary", "Organization Breakdown: Total Activities", "Sum of tracked contribution events for organization users."],
        ["Summary", "Status", "Activity classification label: Active, Moderate, Low, or Inactive."],
        ["Summary", "Definition", "Configured distinct-active-day thresholds; Inactive means a current member with no tracked activity."],
        ["Summary", "Item", "Organization, repository, or member-list scope that could not be fully processed."],
        ["Summary", "Reason", "API or data error recorded for the affected item."],
        ["User Activity", "Organization", "Organization login."],
        ["User Activity", "User", "GitHub login attributed to activity or a current member listed with no activity."],
        ["User Activity", "Activity Status", "Active/Moderate/Low use configured distinct active-day thresholds; Inactive means a current member with no tracked activity."],
        ["User Activity", "Active Days", "Distinct IST calendar dates with at least one tracked event."],
        ["User Activity", "Total Activities", "Sum of event metric columns; a PR may contribute multiple event types."],
        *[["User Activity", label, "Count of this event type attributed to the user during the report period."] for label in METRIC_LABELS],
        ["User Activity", "Average Commits per Active Month", "This person's commits divided by the number of IST calendar months in which they had any tracked activity. A month counts as active even if they made no commits that month; No months means they had no active months."],
        ["User Activity", "Repositories Contributed", "Distinct selected repositories with attributed activity by this user."],
        ["User Activity", "First Activity (IST)", "Earliest tracked event timestamp converted to IST."],
        ["User Activity", "Last Activity (IST)", "Latest tracked event timestamp converted to IST."],
        ["User Activity", "Days Since Last Activity", "Whole elapsed days from report end to last activity; No activity means there was no activity date to compare."],
        ["User Activity", "Org Member", "Yes if listed as a current org member, No if not, Unknown if member lookup failed."],
        ["User Activity", "Teams", "Teams the person belongs to, including memberships inherited through nested teams. Unknown means the team lookup was incomplete; None means no visible team membership."],
        ["User Activity", "Value Streams", "Distinct repository custom-property values for repos where the user had tracked activity; No activity means no stream could be assigned."],
        ["Repository Detail", "Organization", "Organization login that owns the repository."],
        ["Repository Detail", "Repository", "Repository name."],
        ["Repository Detail", "User", "GitHub login attributed to activity in this repository."],
        ["Repository Detail", "Total Activities", "Sum of the user's tracked event metrics in the repository."],
        *[["Repository Detail", label, "Count of this event type attributed to this user in this repository during the report period."] for label in METRIC_LABELS],
        ["Repository Detail", "Active Days", "Distinct IST dates with this user's tracked activity in the repository."],
        ["Repository Detail", "Last Activity (IST)", "Latest tracked event timestamp for this repository/user, converted to IST."],
        ["Management KPIs", "Scope", "Organization login; KPI values are calculated per organization."],
        ["Management KPIs", "Category", "KPI group such as User Adoption, Developer Activity, or Security & Governance."],
        ["Management KPIs", "Metric", "Name of the KPI; the row definition below describes its calculation."],
        ["Management KPIs", "Value", "Calculated result; N/A or Unavailable means data is missing or calculation is undefined."],
        ["Management KPIs", "What this means", "Plain-language explanation of the KPI, where it comes from, and any important caveat."],
        ["Repository Inventory", "Organization", "Organization login that owns the repository."],
        ["Repository Inventory", "Repository", "Selected repository name."],
        ["Repository Inventory", "Value Stream", "Current repository custom-property value(s); semicolon-separated for multiple values."],
        ["Repository Inventory", "Visibility", "Current GitHub visibility: private, internal, public, or Unknown."],
        ["Repository Inventory", "Created (IST)", "Repository creation timestamp converted to IST."],
        ["Repository Inventory", "New in Period", "Yes if repository creation falls within the reporting period."],
        ["Repository Inventory", "Archived", "Yes if GitHub currently marks the repository archived."],
        ["Repository Inventory", "Default Branch", "Repository default branch name; Not set means GitHub returned no default branch."],
        ["Repository Inventory", "Protected Branches", "Protected branch names; None if no protected branches, Unknown if the check failed."],
        ["Repository Inventory", "CODEOWNERS Present", "Yes if a CODEOWNERS file exists at a recognized path on the default branch; syntax is not validated."],
        ["Repository Inventory", "CODEOWNERS Path", "First matching default-branch path among .github/CODEOWNERS, CODEOWNERS, and docs/CODEOWNERS; Not found means none exists at those paths."],
        ["Repository Inventory", "Activity Last 7 Days", "Yes if a tracked event occurred within seven days; Unknown if activity collection failed."],
        ["Repository Inventory", "Total Activities", "Tracked event total during the report period; N/A on activity collection failure."],
        ["Repository Inventory", "Active Contributors", "Distinct linked contributors during the report period; N/A on activity collection failure."],
        *[["Repository Inventory", label, "Count of this event type in the repository during the report period; N/A if activity collection failed."] for label in METRIC_LABELS],
        ["Repository Inventory", "PR Cohort Merge Rate (%)", "PRs opened in the period and merged by cutoff divided by PRs opened in the period, times 100; N/A for an empty cohort."],
        ["Repository Inventory", "Avg First Review Hours", "Mean elapsed hours from PR creation to first non-self submitted review; includes weekends/draft time and excludes unreviewed PRs."],
        ["Repository Inventory", "First Review Sample PRs", "Number of PRs used in the first-review turnaround average."],
        ["Repository Inventory", "Coverage", "Complete if activity collection succeeded without repository errors; otherwise Partial."],
        ["Top Active Repositories", "Organization", "The organization this ranked list belongs to."],
        ["Top Active Repositories", "Rank", "Position within the organization, ordered by activity count; ties are sorted by repository name."],
        ["Top Active Repositories", "Repository", "Repository name. Up to ten repositories are listed for each organization."],
        ["Top Active Repositories", "Value Stream", "Value or values assigned to the repository by its custom property."],
        ["Top Active Repositories", "Total Activities", "Sum of the repository's tracked events during the report period; this is not a productivity score."],
        ["Top Active Repositories", "Active Contributors", "Distinct linked users with tracked activity in the repository during this period."],
        *[["Top Active Repositories", label, f"Number of {label.lower()} recorded in the repository during the report period."] for label in METRIC_LABELS],
        ["Top Active Repositories", "Activity Last 7 Days", "Yes if a tracked event occurred in the last seven days; Unknown if the scan failed."],
        ["Top Active Repositories", "Coverage", "Complete means the activity scan finished; Partial means some repository data may be missing."],
        ["Value Streams", "Value Stream", "Repository custom-property value; multi-select repositories contribute to each value."],
        ["Value Streams", "Organization Names", "Distinct organization logins owning repositories assigned to this stream."],
        ["Value Streams", "Repository Names (Org/Repo)", "Semicolon-separated qualified repository names for quick lookup."],
        ["Value Streams", "Repositories", "Selected repository count for the stream; counts can overlap."],
        ["Value Streams", "Active Org Members", "Distinct current members who contributed to stream repositories; N/A if member lookup failed."],
        ["Value Streams", "Active Contributors", "Distinct linked contributors with activity on stream repositories."],
        ["Value Streams", "Private", "Count of stream-assigned repositories with private visibility."],
        ["Value Streams", "Internal", "Count of stream-assigned repositories with internal visibility."],
        ["Value Streams", "Public", "Count of stream-assigned repositories with public visibility."],
        ["Growth Trends", "Granularity", "Weekly is Monday-start IST week; Monthly is IST calendar month."],
        ["Growth Trends", "Period Start (IST)", "Calendar period start, clipped to the report-window start when the first period is partial."],
        ["Growth Trends", "Period End (IST)", "Calendar period end, clipped to the report end when the last period is partial."],
        ["Growth Trends", "Period Coverage", "Full if the complete period is in-window and ended; otherwise Partial."],
        ["Growth Trends", "Granularity", "Weekly is Monday-start IST week; Monthly is IST calendar month."],
        ["Growth Trends", "Period Start (IST)", "Calendar period start, clipped to the report-window start when the first period is partial."],
        ["Growth Trends", "Period End (IST)", "Calendar period end, clipped to the report end when the last period is partial."],
        ["Growth Trends", "Period Coverage", "Full if the complete period is in-window and ended; otherwise Partial."],
        ["Growth Trends", "Observed User Growth (%)", "Percentage change in distinct observed organization members from the preceding weekly/monthly snapshot. No growth means no adjacent snapshot is available, not a measured 0%; No baseline means the previous snapshot had zero members."],
        ["Self-Hosted Runners", "Organization", "Organization owning the self-hosted runner or repository."],
        ["Self-Hosted Runners", "Runner ID", "GitHub unique numeric runner identifier."],
        ["Self-Hosted Runners", "Runner Name", "Hostname or registered runner identifier."],
        ["Self-Hosted Runners", "Scope", "Organization-wide runner vs repository-specific runner."],
        ["Self-Hosted Runners", "Runner Group", "Runner group assignment used for access control and isolation."],
        ["Self-Hosted Runners", "OS", "Runner operating system reported by the runner agent (Linux, Windows, macOS)."],
        ["Self-Hosted Runners", "Arch", "Processor architecture reported by runner (x64, arm64)."],
        ["Self-Hosted Runners", "Agent Version", "Installed runner agent software version for patch and security compliance."],
        ["Self-Hosted Runners", "Provisioning Type", "Ephemeral (auto-scaling dynamic container/pod) vs Persistent (long-running static VM)."],
        ["Self-Hosted Runners", "Status", "Current runner connectivity status: Online, Offline, or Ephemeral."],
        ["Self-Hosted Runners", "Current State", "Busy if currently executing an active workflow job; Idle if waiting."],
        ["Self-Hosted Runners", "Labels", "Configured runner labels used by workflow runs-on directives."],
        ["Self-Hosted Runners", "Repos Served", "Count of distinct repositories that executed jobs on this runner."],
        ["Self-Hosted Runners", "Total Jobs", "Number of workflow jobs executed on this runner during the report period."],
        ["Self-Hosted Runners", "Successful", "Number of workflow jobs that completed successfully."],
        ["Self-Hosted Runners", "Failed", "Number of workflow jobs that failed during execution."],
        ["Self-Hosted Runners", "Cancelled", "Number of workflow jobs that were cancelled by users."],
        ["Self-Hosted Runners", "Timed Out", "Number of workflow jobs aborted due to runner hang or exceeding timeout limits."],
        ["Self-Hosted Runners", "Success Rate (%)", "Percentage of completed jobs that concluded with success."],
        ["Self-Hosted Runners", "Total Compute (Hours)", "Cumulative execution duration across all jobs executed on this runner."],
        ["Self-Hosted Runners", "Avg Duration (Mins)", "Average job execution duration in minutes."],
        ["Self-Hosted Runners", "Median Duration (Mins)", "50th percentile (median) job execution duration in minutes."],
        ["Self-Hosted Runners", "Min Duration (Mins)", "Minimum observed job execution duration in minutes."],
        ["Self-Hosted Runners", "Max Duration (Mins)", "Maximum observed job execution duration in minutes."],
        ["Self-Hosted Runners", "Avg Queue Wait (s)", "Average wait time from job creation/queuing to runner start execution."],
        ["Self-Hosted Runners", "P95 Queue Wait (s)", "95th percentile queue wait time (peak scheduling SLA delay)."],
        ["Self-Hosted Runners", "SLA Breaches (>60s)", "Count of workflow jobs delayed more than 60 seconds in queue waiting for a runner."],
        ["Self-Hosted Runners", "First Active (IST)", "Timestamp of the earliest job executed on this runner during the lookback period."],
        ["Self-Hosted Runners", "Last Active (IST)", "Timestamp of the most recent job completed or started on this runner."],
        ["Metric Notes", "Sheet / Field / Definition", "Workbook sheet, field or metric, and its meaning or calculation."],
        ["All Sheets", "Scope", "Only enabled repositories listed in repos.csv and accessible to the GitHub App; not an enterprise-wide census."],
        ["All Sheets", "Unmapped / Unknown / N/A", "Unmapped: property absent/empty. Unknown: lookup failed. N/A: data unavailable or calculation undefined. Empty cells use short, field-specific messages such as No activity or No snapshot."],
        ["All Sheets", "Partial Run", "Failures can make totals lower bounds; consult Items Not Processed and Coverage."],
        ["All Sheets", "Tracked Activity", "Default-branch commits, PR opens/merges, submitted reviews, review comments, issue opens, and issue/PR comments; excludes logins, browsing, cloning, and non-default-branch pushes."],
        ["All Sheets", "Activity Total", "Sum of tracked event types; not a productivity or quality score."],
        ["All Sheets", "Adoption History", "Snapshot growth requires retained history and repeated complete runs. Exact organization join counts require a complete audit-log CSV; sign-in events are not reported."],
    ])
    seen_kpis: Set[Tuple[str, str]] = set()
    for _, category, metric, _, definition in kpis:
        if (category, metric) not in seen_kpis:
            notes.append(["Management KPIs", f"{category}: {metric}", definition])
            seen_kpis.add((category, metric))
    tables["Metric Notes"] = (("Sheet", "Field / Metric", "What this means"), notes)
    return tables


# --------------------------------------------------------------------------- output

EMPTY_CELL_MESSAGES = {
    "Average Commits per Active Month": "No months",
    "First Activity (IST)": "No activity",
    "Last Activity (IST)": "No activity",
    "Days Since Last Activity": "No activity",
    "Value Streams": "No activity",
    "Default Branch": "Not set",
    "CODEOWNERS Path": "Not found",
    "Snapshot (IST)": "No snapshot",
    "First Active (IST)": "No activity",
    "Last Active (IST)": "No activity",
    "Avg Queue Wait (s)": "N/A",
    "P95 Queue Wait (s)": "N/A",
    "Avg Duration (Mins)": "N/A",
    "Median Duration (Mins)": "N/A",
    "Min Duration (Mins)": "N/A",
    "Max Duration (Mins)": "N/A",
    "Success Rate (%)": "N/A",
    "Availability (%)": "N/A",
}


def safe_cell(value: Any, header: Optional[str] = None) -> Any:
    """Neutralise spreadsheet formulas and explain expected empty report cells."""
    if (value is None or value == "") and header in EMPTY_CELL_MESSAGES:
        value = EMPTY_CELL_MESSAGES[header]
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return f"'{value}"
    return value


def write_csv(path: Path, headers: Tuple[str, ...], rows: List[list]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows([safe_cell(value, headers[index]) for index, value in enumerate(row)] for row in rows)


def write_table(ws, top: int, headers, rows: List[list], status_col: Optional[int] = None,
                bar_col: Optional[int] = None) -> int:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=top, column=col, value=header)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    center_headers = {
        "Value", "Status", "Current State", "Scope", "OS", "Arch", "Agent Version",
        "Provisioning Type", "Coverage", "CODEOWNERS Present", "Protected Branches",
        "Availability (%)", "Success Rate (%)", "Total Jobs", "Total Jobs Run", "Total Compute (Hours)",
        "Avg Duration (Mins)", "Median Duration (Mins)", "Min Duration (Mins)", "Max Duration (Mins)",
        "Avg Queue Wait (s)", "P95 Queue Wait (s)", "SLA Breaches (>60s)", "SLA Breach Rate (%)",
        "Timed Out", "Successful", "Failed", "Cancelled", "Runner ID", "Repos Served", "Active Repos",
        "Org-Level", "Repo-Level", "Online", "Offline", "Currently Busy", "Ephemeral Count",
        "First Active (IST)", "Last Active (IST)", "Last Active"
    }

    for row_idx, row in enumerate(rows, start=top + 1):
        banded = (row_idx - top) % 2 == 0
        for col_idx, value in enumerate(row, start=1):
            header = headers[col_idx - 1] if col_idx - 1 < len(headers) else ""
            cell = ws.cell(row=row_idx, column=col_idx,
                           value=safe_cell(value, header))
            cell.border = BORDER
            if banded:
                cell.fill = BAND_FILL
            if (isinstance(value, (int, float))
                    or header in center_headers
                    or header.endswith("(%)") or header.endswith("(s)")
                    or header.endswith("(Hours)") or header.endswith("(Mins)")
                    or str(value).strip() in {"N/A", "None", "Yes", "No", "Unknown", "Complete", "Partial", "Idle", "Busy"}
                    or re.match(r"^\d+\s*\(\d+\s+online\)$", str(value).strip())):
                cell.alignment = CENTER
        if status_col is not None and row[status_col] in STATUS_STYLES:
            fill, font = STATUS_STYLES[row[status_col]]
            cell = ws.cell(row=row_idx, column=status_col + 1)
            cell.fill, cell.font, cell.alignment = PatternFill("solid", fgColor=fill), Font(bold=True, color=font), CENTER

    last_row = top + len(rows)
    if bar_col is not None and rows:
        letter = get_column_letter(bar_col + 1)
        ws.conditional_formatting.add(f"{letter}{top + 1}:{letter}{last_row}",
                                      DataBarRule(start_type="min", end_type="max", color="5B9BD5"))
    return last_row


def autofit(ws, skip_rows: int = 0) -> None:
    widths: Dict[int, int] = defaultdict(int)
    for row in ws.iter_rows(min_row=skip_rows + 1):
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths[cell.column], len(str(cell.value)))
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 3, 12), 50)


def finalize_sheet(ws, header_row: int, last_row: int, columns: int, freeze: str) -> None:
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(columns)}{max(last_row, header_row)}"
    ws.freeze_panes = freeze
    ws.row_dimensions[header_row].height = 32
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    autofit(ws)


def section_title(ws, row: int, title: str) -> int:
    ws.cell(row=row, column=1, value=title).font = Font(bold=True, size=12, color=NAVY)
    return row + 1



def write_workbook(path: Path, cfg: Config, since: datetime, end: datetime, user_rows: List[list],
                   repo_rows: List[list], org_rows: List[list], failures: List[Tuple[str, str]],
                   repos_scanned: int,
                   management_tables: Optional[Dict[str, Tuple[Tuple[str, ...], List[list]]]] = None,
                   runner_rows: Optional[List[list]] = None,
                   org_runner_rows: Optional[List[list]] = None) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "GitHub User Activity Report"
    summary["A1"].font = Font(bold=True, size=16, color=NAVY)
    since_ist, end_ist = since.astimezone(IST), end.astimezone(IST)
    summary["A2"] = (f"Reporting period: {since_ist:{DATE_FMT}} to {end_ist:{DATE_FMT}} ({cfg.lookback_days} days)"
                     f"  |  Generated: {end_ist:{DATE_FMT} %H:%M} IST")
    summary["A2"].font = Font(italic=True, color="595959")

    status_counts = Counter(row[STATUS_IDX] for row in user_rows)
    status_idx = RUNNER_HEADERS.index("Status") if runner_rows and len(runner_rows[0]) > RUNNER_HEADERS.index("Status") else 8
    compute_idx = RUNNER_HEADERS.index("Total Compute (Hours)") if runner_rows and len(runner_rows[0]) > RUNNER_HEADERS.index("Total Compute (Hours)") else 12
    total_runners_count = len(runner_rows) if runner_rows else 0
    online_runners_count = sum(1 for r in (runner_rows or []) if len(r) > status_idx and str(r[status_idx]).capitalize() == "Online")
    total_compute_hours = sum(r[compute_idx] for r in (runner_rows or []) if len(r) > compute_idx and isinstance(r[compute_idx], (int, float)))

    kpis = [
        ["Organizations", len(org_rows)],
        ["Repositories Scanned", repos_scanned],
        ["Items Not Processed", len(failures)],
        ["Users Reported", len(user_rows)],
        *([f"{status} Users", status_counts[status]] for status in STATUSES),
        ["Total Activities", sum(row[TOTAL_IDX] for row in user_rows)],
        *([label, sum(row[METRIC_START_IDX + i] for row in user_rows)] for i, label in enumerate(METRIC_LABELS)),
        ["Self-Hosted Runners", f"{total_runners_count} ({online_runners_count} online)" if total_runners_count else "0"],
        ["Self-Hosted Compute (Hours)", round(total_compute_hours, 2)],
    ]
    row = write_table(summary, 4, ("Metric", "Value"), kpis) + 2

    row = section_title(summary, row, "Organization Breakdown")
    row = write_table(summary, row, ORG_HEADERS, org_rows, bar_col=len(ORG_HEADERS) - 1) + 2

    row = section_title(summary, row, "Status Definitions (distinct days with activity in period)")
    definitions = [
        ["Active", f"Activity on {cfg.active_min_days} or more days"],
        ["Moderate", f"Activity on {cfg.moderate_min_days} to {cfg.active_min_days - 1} days"],
        ["Low", f"Activity on 1 to {cfg.moderate_min_days - 1} days"],
        ["Inactive", "Organization member with no recorded activity"],
    ]
    row = write_table(summary, row, ("Status", "Definition"), definitions, status_col=0) + 2

    if failures:
        row = section_title(summary, row, "Items Not Processed")
        write_table(summary, row, ("Item", "Reason"), [list(item) for item in failures])
    autofit(summary, skip_rows=2)
    summary.sheet_view.showGridLines = False

    users_ws = wb.create_sheet("User Activity")
    last = write_table(users_ws, 1, USER_HEADERS, user_rows, status_col=STATUS_IDX, bar_col=TOTAL_IDX)
    finalize_sheet(users_ws, 1, last, len(USER_HEADERS), "C2")
    stream_column = USER_HEADERS.index("Value Streams") + 1
    users_ws.column_dimensions[get_column_letter(stream_column)].width = 28
    for row_idx in range(2, last + 1):
        users_ws.cell(row_idx, stream_column).alignment = Alignment(vertical="top", wrap_text=True)
        users_ws.row_dimensions[row_idx].height = max(users_ws.row_dimensions[row_idx].height or 15, 36)

    repo_ws = wb.create_sheet("Repository Detail")
    last = write_table(repo_ws, 1, REPO_HEADERS, repo_rows, bar_col=REPO_HEADERS.index("Total Activities"))
    finalize_sheet(repo_ws, 1, last, len(REPO_HEADERS), "D2")

    def _create_runner_sheet():
        runner_ws = wb.create_sheet("Self-Hosted Runners")
        curr_row = 1
        curr_row = section_title(runner_ws, curr_row, "Actions Self-Hosted Runner Fleet & Performance")
        if runner_rows:
            last = write_table(runner_ws, curr_row, RUNNER_HEADERS, runner_rows)
            finalize_sheet(runner_ws, curr_row, last, len(RUNNER_HEADERS), "C3")
            runner_ws.sheet_view.showGridLines = False

            status_col_idx = RUNNER_HEADERS.index("Status") + 1
            for r_idx in range(curr_row + 1, last + 1):
                cell = runner_ws.cell(r_idx, status_col_idx)
                val = str(cell.value).capitalize()
                if val == "Online":
                    cell.fill = PatternFill("solid", fgColor="C6EFCE")
                    cell.font = Font(bold=True, color="006100")
                elif val in ("Offline", "Ephemeral"):
                    cell.fill = PatternFill("solid", fgColor="FFC7CE")
                    cell.font = Font(bold=True, color="9C0006")
            curr_row = last + 3
        else:
            empty_row = [["No self-hosted runners found (verify GitHub App permission: Organization self-hosted runners: Read-only)"]
                         + ["-"] * (len(RUNNER_HEADERS) - 1)]
            last = write_table(runner_ws, curr_row, RUNNER_HEADERS, empty_row)
            finalize_sheet(runner_ws, curr_row, last, len(RUNNER_HEADERS), "C3")
            runner_ws.sheet_view.showGridLines = False
            curr_row = last + 3

        if org_runner_rows:
            curr_row = section_title(runner_ws, curr_row, "Organization Runner Capacity & SLA Summary")
            write_table(runner_ws, curr_row, ORG_RUNNER_HEADERS, org_runner_rows)
            autofit(runner_ws)
        else:
            curr_row = section_title(runner_ws, curr_row, "Organization Runner Capacity & SLA Summary")
            write_table(runner_ws, curr_row, ORG_RUNNER_HEADERS, [["Default", 0, 0, 0, 0, 0, "N/A", 0, 0, 0, 0, 0, 0, 0, 0, "N/A", 0.0, "N/A", "N/A", 0, "N/A"]])
            autofit(runner_ws)

    management_tables = management_tables or {}
    management_sheets = {}
    runner_sheet_created = False
    for title, (headers, rows) in management_tables.items():
        if title in {"Management KPIs", "Metric Notes"} and not runner_sheet_created and cfg.include_runner_metrics:
            _create_runner_sheet()
            runner_sheet_created = True

        sheet = wb.create_sheet(title)
        management_sheets[title] = sheet
        last = write_table(sheet, 1, headers, rows)
        finalize_sheet(sheet, 1, last, len(headers), "B2")
        sheet.sheet_view.showGridLines = False
        if title == "Repository Inventory":
            colors = {"Yes": STATUS_STYLES["Active"], "No": STATUS_STYLES["Inactive"],
                      "None": STATUS_STYLES["Inactive"],
                      "Unknown": STATUS_STYLES["Moderate"], "Complete": STATUS_STYLES["Active"],
                      "Partial": STATUS_STYLES["Moderate"]}
            for column in ("Protected Branches", "CODEOWNERS Present", "Coverage"):
                for cells in sheet.iter_rows(min_row=2, min_col=headers.index(column) + 1,
                                             max_col=headers.index(column) + 1):
                    cell = cells[0]
                    status = cell.value if column != "Protected Branches" or cell.value in {"None", "Unknown"} else "Yes"
                    fill, font = colors[status]
                    cell.fill = PatternFill("solid", fgColor=fill)
                    cell.font = Font(bold=True, color=font)
            sheet.column_dimensions[get_column_letter(headers.index("Protected Branches") + 1)].width = 40
        if title == "Top Active Repositories" and rows:
            column = get_column_letter(headers.index("Total Activities") + 1)
            sheet.conditional_formatting.add(
                f"{column}2:{column}{last}",
                DataBarRule(start_type="min", end_type="max", color="5B9BD5"))
        if title in {"Management KPIs", "Metric Notes"}:
            letter = get_column_letter(len(headers))
            sheet.column_dimensions[letter].width = 100
            val_col_idx = headers.index("Value") + 1 if "Value" in headers else None
            for cells in sheet.iter_rows(min_row=2):
                for col_idx, cell in enumerate(cells, start=1):
                    h_align = "center" if col_idx == val_col_idx else None
                    cell.alignment = Alignment(horizontal=h_align, vertical="center", wrap_text=True)

    if not runner_sheet_created and cfg.include_runner_metrics:
        _create_runner_sheet()

    wb.save(path)


# --------------------------------------------------------------------------- entry point

class ColorFormatter(logging.Formatter):
    COLORS = {
        logging.DEBUG: "\033[90m",
        logging.INFO: "\033[32m",
        logging.WARNING: "\033[33m",
        logging.ERROR: "\033[31m",
        logging.CRITICAL: "\033[1;31m",
    }
    RESET = "\033[0m"

    def format(self, record: logging.LogRecord) -> str:
        return f"{self.COLORS.get(record.levelno, '')}{super().format(record)}{self.RESET}"


def console_supports_color() -> bool:
    if os.environ.get("NO_COLOR") or not sys.stderr.isatty():
        return False
    if os.name != "nt":
        return True
    # Enable ANSI escape processing on Windows consoles.
    kernel32 = ctypes.windll.kernel32
    handle = kernel32.GetStdHandle(-12)
    mode = ctypes.c_uint32()
    return bool(kernel32.GetConsoleMode(handle, ctypes.byref(mode))
                and kernel32.SetConsoleMode(handle, mode.value | 0x0004))


def setup_logging(level: str, log_file: Path) -> None:
    console_fmt = "%(asctime)s | %(levelname)-8s | %(message)s"
    console = logging.StreamHandler()
    console.setFormatter(ColorFormatter(console_fmt) if console_supports_color() else logging.Formatter(console_fmt))
    file_handler = logging.FileHandler(log_file, encoding="utf-8")
    file_handler.setFormatter(logging.Formatter("%(asctime)s | %(levelname)-8s | %(threadName)s | %(message)s"))
    LOG.addHandler(console)
    LOG.addHandler(file_handler)
    LOG.setLevel(level)
    logging.getLogger("urllib3").setLevel(logging.WARNING)


def run(cfg: Config, end: datetime, stamp: str) -> int:
    entries = load_repo_entries(cfg.repos_csv)
    history_file = cfg.history_file or cfg.output_dir / "adoption_history.json"
    history = load_history(history_file)
    scope = history_scope(cfg, entries)
    since = end - timedelta(days=cfg.lookback_days)
    LOG.info("Reporting window %s -> %s IST (%d days)", f"{since.astimezone(IST):{DATE_FMT}}",
             f"{end.astimezone(IST):{DATE_FMT}}", cfg.lookback_days)
    audit_joined_users = (load_onboarded_users(cfg.audit_log_csv, since, end)
                          if cfg.audit_log_csv else None)
    if audit_joined_users is not None:
        LOG.info("Loaded audit-log join events for %d organization(s)", len(audit_joined_users))

    client = GitHubClient(cfg)
    targets, failures = resolve_targets(client, entries, cfg)
    if not targets:
        LOG.error("No repositories to scan")
        return 1

    reports: Dict[Tuple[str, str], RepositoryReport] = {}
    with ThreadPoolExecutor(max_workers=cfg.max_workers, thread_name_prefix="worker") as pool:
        futures = {pool.submit(collect_repository, client, org, repo, since, end, cfg.exclude_bots,
                       cfg.value_stream_property): (org, repo)
                   for org, repo in targets}
        for done, future in enumerate(as_completed(futures), start=1):
            org, repo = futures[future]
            try:
                collected = future.result()
                reports[(org, repo)] = collected
                for error in collected.errors:
                    failures.append((f"{org}/{repo}", error))
                    LOG.warning("%s/%s: %s", org, repo, error)
                LOG.info("[%d/%d] %s/%s: %d contributors", done, len(targets), org, repo,
                         len(collected.users))
            except (GitHubAPIError, requests.RequestException, KeyError, ValueError) as exc:
                LOG.error("[%d/%d] %s/%s failed: %s", done, len(targets), org, repo, exc)
                failures.append((f"{org}/{repo}", str(exc)))

    orgs = sorted({org for org, _ in targets}, key=str.lower)
    members = fetch_members(client, orgs, cfg.exclude_bots, failures)
    teams_by_user, team_streams_by_user, org_team_complete = fetch_team_access(
        client, orgs, reports, cfg.exclude_bots, failures, cfg.value_stream_property)

    runners: List[RunnerInfo] = []
    runner_rows: List[list] = []
    org_runner_rows: List[list] = []
    org_runner_metrics: Dict[str, RunnerJobMetrics] = {}
    if cfg.include_runner_metrics:
        LOG.info("Collecting Actions self-hosted runner metrics across %d organization(s)...", len(orgs))
        runners, runner_metrics, org_runner_metrics = collect_runners_and_metrics(
            client, orgs, targets, since, end, cfg, failures
        )
        runner_rows, org_runner_rows = build_runner_rows(runners, runner_metrics, org_runner_metrics, orgs)
        LOG.info("Actions runner metrics: %d self-hosted runner(s) processed", len(runners))

    target_counts = Counter(org for org, _ in targets)
    failed_orgs = {
        org for org in target_counts
        if any(item.lower() == org.lower() or item.lower().startswith((f"{org}/".lower(), f"{org} (".lower()))
               for item, _ in failures)
    }
    organization_status = {
        org: (len([key for key in reports if key[0] == org]) == target_counts[org]
              and org in members and org not in failed_orgs
              and org_team_complete.get(org, False)
              and all(not result.errors for (owner, _), result in reports.items() if owner == org))
        for org in orgs
    }
    complete = (not failures and len(reports) == len(targets) and len(members) == len(orgs)
                and all(org_team_complete.get(org, False) for org in orgs))
    snapshots = [snapshot for snapshot in history["snapshots"]
                 if snapshot["scope"] == scope and parse_ts(snapshot["at"]) < end]
    if complete:
        snapshot = membership_snapshot(scope, end, members, reports)
        snapshots.append(snapshot)
        history["snapshots"] = [item for item in history["snapshots"]
                                if not (item["scope"] == scope and item["at"] == snapshot["at"])] + [snapshot]
    management_tables = build_management_tables(
        reports, members, since, end, snapshots, organization_status, audit_joined_users,
        runners=runners, org_runner_metrics=org_runner_metrics)

    results = {key: result.users for key, result in reports.items() if result.activity_complete}
    user_streams = build_user_streams(reports, since, end, team_streams_by_user)
    user_rows, repo_rows = build_rows(results, members, cfg, end, user_streams, teams_by_user, org_team_complete)
    org_rows = build_org_rows(user_rows, results)

    csv_path = cfg.output_dir / f"{cfg.file_prefix}_{stamp}.csv"
    xlsx_path = cfg.output_dir / f"{cfg.file_prefix}_{stamp}.xlsx"
    write_csv(csv_path, USER_HEADERS, user_rows)
    write_workbook(xlsx_path, cfg, since, end, user_rows, repo_rows, org_rows, failures, len(results),
                   management_tables, runner_rows=runner_rows, org_runner_rows=org_runner_rows)
    if complete:
        save_history(history_file, history)
    else:
        LOG.warning("Partial run: adoption history was not updated; management metrics may be incomplete")

    LOG.info("CSV report : %s", csv_path)
    LOG.info("XLSX report: %s", xlsx_path)
    if failures:
        LOG.warning("Completed with %d item(s) not processed - see Summary sheet", len(failures))
        return 2
    LOG.info("Completed successfully: %d users across %d repositories", len(user_rows), len(results))
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="GitHub user activity report via GitHub App.")
    parser.add_argument("--config", default="config.json", type=Path, help="Path to config JSON (default: config.json)")
    args = parser.parse_args()

    end = datetime.now(timezone.utc)
    stamp = end.strftime("%Y%m%d_%H%M%S")
    try:
        cfg = load_config(args.config)
        cfg.output_dir.mkdir(parents=True, exist_ok=True)
        setup_logging(cfg.log_level, cfg.output_dir / f"{cfg.file_prefix}_{stamp}.log")
        return run(cfg, end, stamp)
    except ConfigError as exc:
        print(f"Configuration error: {exc}", file=sys.stderr)
        return 1
    except (GitHubAPIError, requests.RequestException, OSError) as exc:
        LOG.critical("Fatal error: %s", exc)
        print(f"Fatal error: {exc}", file=sys.stderr)
        return 1
    except KeyboardInterrupt:
        print("Interrupted.", file=sys.stderr)
        return 130

if __name__ == "__main__":
    sys.exit(main())
